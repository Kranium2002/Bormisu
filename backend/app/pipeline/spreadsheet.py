"""Spreadsheet path (BACKEND.md 7.10): .xlsx and .csv uploads become boreholes without page images or OCR.

app.pipeline.orchestrator runs these stages for spreadsheet documents instead of triage / classify / OCR /
extract:

sheet_read
    Opens the file without a model: every sheet with openpyxl (stored values and cached formula results,
    merged ranges, notes) or the CSV with the csv module (encoding and delimiter sniffed). The document gets
    page_count null and no pages.
sheet_map
    Finds the table regions (header rows, unit rows) and sends the extractor model a preview only: per sheet
    layout the rows above the table header, the header and unit rows, a few data rows, the notes and the
    distinct values of short text columns; never a whole table. The model answers a DECLARATIVE mapping in a
    strict schema (SheetMapping): per region the header row, column -> target field from the whitelist in
    data/spreadsheet.yaml, the unit printed per column, the locale, row-type rules and the borehole-id
    source. The answer is stored in extraction_runs under sha256(file sha256, model, prompt and schema
    versions, hash of what the model reads), so a rerun never pays again.
sheet_extract
    Validates the mapping against the workbook (every referenced sheet, row, column and cell must exist, a
    mapped column must print a title, a unit is used only where it is printed), applies it to EVERY row in
    Python and builds one record per borehole id. The rows of a borehole become an ExtractionResult whose
    "blocks" are cells, and app.pipeline.finalize.build_record normalises it exactly like the PDF path
    (app.normalize: units, the file's decimal style, SPT, DIN codes, dates, coordinates, a stated or an
    inferred coordinate system). Every provenance entry then points at its cell ("Sheet!A1"; plain "A1" for a
    CSV) with page and box null and method read (derived or guessed where Python computed or assumed a
    value). Boreholes are created through app.review.service.create_borehole, one transaction each and
    idempotent per borehole id; app.pipeline.finalize.finalize_document then asks the coordinate-system
    questions and sets the document status.

Nothing here depends on a file name, sheet name, project, label, language or layout: vocabularies,
thresholds and the prompt live in backend/data. Unknown stays null: a value is never invented, a unit is
used only where it is printed and a formula without a stored result is never computed.
"""

from __future__ import annotations

import asyncio
import csv
import hashlib
import io
import itertools
import json
import logging
import math
import re
import zipfile
from collections import Counter
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass, field, replace
from datetime import UTC, date, datetime, time
from functools import lru_cache
from pathlib import Path
from typing import Any, Literal, Protocol, cast, get_args

from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy import delete, func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from app.clarify.service import clarification_out, create_clarification
from app.config import get_settings
from app.db.models import Borehole, Clarification, Document, ExtractionRun, Page
from app.db.session import get_sessionmaker
from app.events.writer import emit
from app.ids import new_id
from app.normalize import numbers, units
from app.normalize.lithology import valid_din_code
from app.pipeline.finalize import (
    DEPTH_UNIT_OPTIONS,
    GuessedDepth,
    RecordDraft,
    build_record,
    finalize_document,
    length_unit_in,
)
from app.pipeline.orchestrator import StageError
from app.pipeline.schemas import (
    ExtractedCoreRun,
    ExtractedGroundwater,
    ExtractedHeader,
    ExtractedLayer,
    ExtractedSample,
    ExtractedSpt,
    ExtractedTest,
    ExtractionResult,
    RawCoordinates,
    RawValue,
)
from app.pipeline.types import BlockRef, ExtractionOutcome, PageGroup
from app.providers.base import ProviderError, token_cost
from app.providers.llm_openrouter import text_part
from app.review.paths import PathError, get_value, set_value
from app.review.service import canonical_label, create_borehole
from app.schemas.boreholes import ProvenanceIn
from app.services.documents import emit_document_updated
from app.storage import get_storage

logger = logging.getLogger("geomemory.pipeline.spreadsheet")

DATA_FILE = "spreadsheet.yaml"
PROMPT_FILE = "sheet_mapping.txt"
PROMPT_VERSION = "1"
# 2: targets borehole_name, project_id, borehole_length, engineering_soil_unit and geologic_age.
MAPPING_SCHEMA_VERSION = "2"
MAPPING_GROUP = "sheet-mapping"
RUN_DONE = "done"
ACTOR = "pipeline"
CHARS_PER_TOKEN = 3.5
# The sheet name a CSV gets in the preview; its cells are cited without a sheet ("B5").
CSV_SHEET = "csv"
# Group keys of spreadsheet boreholes: one per canonical borehole id of the document.
GROUP_PREFIX = "sheet:"
EXTENSIONS = frozenset({".xlsx", ".csv"})

RowTypeLiteral = Literal["layer", "spt", "core_run", "groundwater", "sample", "test", "borehole", "ignore"]
GroundwaterTypeLiteral = Literal["strike", "rest", "standpipe", "other"]
TargetKind = Literal["text", "date", "length", "percent", "angle", "coordinate", "test"]
TargetScope = Literal["borehole", "row", "any"]

# Targets this module knows how to read; data/spreadsheet.yaml may list only these (checked on load).
TARGET_SUPPORT = frozenset(
    {
        "borehole_id",
        "borehole_name",
        "project_name",
        "project_id",
        "client",
        "location_name",
        "drilling_contractor",
        "drilling_method",
        "start_date",
        "end_date",
        "x",
        "y",
        "coordinate_system",
        "ground_level",
        "elevation_datum",
        "total_depth",
        "borehole_length",
        "inclination",
        "notes",
        "depth_from",
        "depth_to",
        "description",
        "din4023_code",
        "uscs",
        "colour",
        "consistency",
        "moisture",
        "weathering",
        "stratigraphic_name",
        "engineering_soil_unit",
        "geologic_age",
        "spt_depth",
        "spt_blows",
        "spt_n",
        "spt_increment_note",
        "rqd",
        "recovery",
        "groundwater_depth",
        "groundwater_date",
        "sample_name",
        "sample_type",
        "sample_from",
        "sample_to",
        "test_value",
        "test_depth",
        "test_sample",
    }
)
# Targets several columns of one region may map to (descriptions are joined; one test per column).
MULTI_COLUMN_TARGETS = frozenset({"description", "test_value"})
# Dimension of the unit a column of each kind may print (data/units.yaml); None: no unit is read.
_KIND_DIMENSION: dict[str, str | None] = {
    "text": None,
    "date": None,
    "coordinate": None,  # never converted: the unit belongs to the coordinate system
    "length": "length",
    "percent": "ratio",
    "angle": "angle",
}


class SpreadsheetError(RuntimeError):
    """The file cannot be read as a spreadsheet. The message is safe to show."""


# ----------------------------------------------------------------------------------------------
# Configuration (data/spreadsheet.yaml)
# ----------------------------------------------------------------------------------------------


@dataclass(frozen=True)
class TargetSpec:
    name: str
    scope: TargetScope
    kind: TargetKind
    description: str


@dataclass(frozen=True)
class SheetConfig:
    csv_encodings: tuple[str, ...]
    csv_delimiters: tuple[str, ...]
    csv_sniff_lines: int
    max_cells: int
    max_uncompressed_bytes: int
    scan_rows: int
    min_header_cells: int
    min_text_share: float
    check_rows: int
    unit_cell_max_chars: int
    rows_above_header: int
    rows_without_table: int
    sample_rows: int
    max_columns: int
    max_cell_chars: int
    max_distinct_values: int
    max_distinct_chars: int
    max_notes: int
    max_layouts: int
    max_further_tables: int
    tokens_floor: int
    tokens_per_layout: int
    tokens_cap: int
    text_join: str
    header_repeat_share: float
    row_types: dict[str, str]
    targets: dict[str, TargetSpec]


def _section(data: Mapping[str, Any], name: str) -> Mapping[str, Any]:
    value = data.get(name)
    if not isinstance(value, Mapping):
        raise ValueError(f"{DATA_FILE}: section {name!r} is missing")
    return value


@lru_cache(maxsize=1)
def config() -> SheetConfig:
    data = numbers.load_data(DATA_FILE)
    reader, detection = _section(data, "reader"), _section(data, "detection")
    preview, mapping = _section(data, "preview"), _section(data, "mapping")
    row_types = {str(k): str(v) for k, v in _section(data, "row_types").items()}
    if set(row_types) != set(get_args(RowTypeLiteral)):
        raise ValueError(f"{DATA_FILE}: row_types must be exactly {', '.join(get_args(RowTypeLiteral))}")
    targets: dict[str, TargetSpec] = {}
    for name, entry in _section(data, "targets").items():
        key = str(name)
        if key not in TARGET_SUPPORT:
            raise ValueError(f"{DATA_FILE}: target {key!r} is not supported by the spreadsheet reader")
        scope, kind = str(entry.get("scope")), str(entry.get("kind"))
        if scope not in get_args(TargetScope) or kind not in get_args(TargetKind):
            raise ValueError(f"{DATA_FILE}: target {key!r} has an unknown scope or kind")
        targets[key] = TargetSpec(
            key, cast(TargetScope, scope), cast(TargetKind, kind), str(entry.get("description") or "")
        )
    if "borehole_id" not in targets:
        raise ValueError(f"{DATA_FILE}: the borehole_id target is required")
    return SheetConfig(
        csv_encodings=tuple(str(e) for e in reader["csv_encodings"]),
        csv_delimiters=tuple(str(d) for d in reader["csv_delimiters"]),
        csv_sniff_lines=int(reader["csv_sniff_lines"]),
        max_cells=int(reader["max_cells"]),
        max_uncompressed_bytes=int(float(reader["max_uncompressed_mb"]) * 1024 * 1024),
        scan_rows=int(detection["scan_rows"]),
        min_header_cells=int(detection["min_header_cells"]),
        min_text_share=float(detection["min_text_share"]),
        check_rows=int(detection["check_rows"]),
        unit_cell_max_chars=int(detection["unit_cell_max_chars"]),
        rows_above_header=int(preview["rows_above_header"]),
        rows_without_table=int(preview["rows_without_table"]),
        sample_rows=int(preview["sample_rows"]),
        max_columns=int(preview["max_columns"]),
        max_cell_chars=int(preview["max_cell_chars"]),
        max_distinct_values=int(preview["max_distinct_values"]),
        max_distinct_chars=int(preview["max_distinct_chars"]),
        max_notes=int(preview["max_notes"]),
        max_layouts=int(preview["max_layouts"]),
        max_further_tables=int(preview["max_further_tables"]),
        tokens_floor=int(mapping["output_tokens_floor"]),
        tokens_per_layout=int(mapping["output_tokens_per_layout"]),
        tokens_cap=int(mapping["output_tokens_cap"]),
        text_join=str(mapping["text_join"]),
        header_repeat_share=float(mapping["header_repeat_share"]),
        row_types=row_types,
        targets=targets,
    )


# ----------------------------------------------------------------------------------------------
# Workbook model
# ----------------------------------------------------------------------------------------------

CellKind = Literal["text", "number", "date", "bool"]
Pos = tuple[int, int]  # (row, column), 1-based

_REF = re.compile(r"^\$?([A-Za-z]{1,3})\$?([1-9][0-9]{0,6})$")
_COLUMN = re.compile(r"^\$?([A-Za-z]{1,3})$")
_PLAIN_SHEET = re.compile(r"^[A-Za-z_][A-Za-z0-9_.]*$")
MAX_COLUMN = 16384  # XFD


def column_letter(col: int) -> str:
    letters = ""
    while col > 0:
        col, rest = divmod(col - 1, 26)
        letters = chr(65 + rest) + letters
    return letters


def column_index(text: str | None) -> int | None:
    """1-based index of a column letter ("C" -> 3); None for anything else."""
    match = _COLUMN.match((text or "").strip())
    if match is None:
        return None
    index = 0
    for ch in match.group(1).upper():
        index = index * 26 + ord(ch) - 64
    return index if 1 <= index <= MAX_COLUMN else None


def parse_ref(text: str | None) -> Pos | None:
    """(row, column) of an A1 reference ("B12" -> (12, 2)); None for anything else."""
    match = _REF.match((text or "").strip())
    if match is None:
        return None
    col = column_index(match.group(1))
    return None if col is None else (int(match.group(2)), col)


def a1(row: int, col: int) -> str:
    return f"{column_letter(col)}{row}"


def _norm(text: str | None) -> str:
    return " ".join((text or "").split()).casefold()


@dataclass(frozen=True)
class SheetCell:
    """One cell with a value. A cell inside a merged range repeats its anchor's value; `source` is where
    the value is printed (the anchor), so provenance always points at the printed cell."""

    row: int
    col: int
    text: str  # the stored value as text (provenance raw)
    parse: str  # what the number parser reads (stored numbers in a form every decimal style reads alike)
    kind: CellKind
    source: Pos
    formula: str | None = None

    @property
    def merged_copy(self) -> bool:
        return self.source != (self.row, self.col)


@dataclass
class Sheet:
    name: str
    index: int
    cells: dict[Pos, SheetCell]
    pending: dict[Pos, str] = field(default_factory=dict)  # formulas without a stored result, error values
    notes: dict[Pos, str] = field(default_factory=dict)
    merged: tuple[str, ...] = ()
    merged_anchors: frozenset[Pos] = frozenset()
    max_row: int = 0
    max_col: int = 0
    _rows: dict[int, list[SheetCell]] | None = None

    def cell(self, row: int, col: int) -> SheetCell | None:
        return self.cells.get((row, col))

    def _index(self) -> dict[int, list[SheetCell]]:
        if self._rows is None:
            rows: dict[int, list[SheetCell]] = {}
            for cell in self.cells.values():
                rows.setdefault(cell.row, []).append(cell)
            for items in rows.values():
                items.sort(key=lambda c: c.col)
            self._rows = rows
        return self._rows

    def row_cells(self, row: int) -> list[SheetCell]:
        return self._index().get(row, [])

    def content_rows(self) -> list[int]:
        return sorted(self._index())


@dataclass
class Workbook:
    kind: Literal["xlsx", "csv"]
    sheets: list[Sheet]

    def sheet(self, name: str) -> Sheet | None:
        for sheet in self.sheets:
            if sheet.name == name:
                return sheet
        folded = [s for s in self.sheets if s.name.strip().casefold() == name.strip().casefold()]
        return folded[0] if len(folded) == 1 else None

    def ref(self, sheet: Sheet, pos: Pos) -> str:
        """Provenance cell: "Sheet!A1" ('quoted' when the name needs it); plain "A1" for a CSV."""
        cell = a1(*pos)
        if self.kind == "csv":
            return cell
        name = sheet.name if _PLAIN_SHEET.match(sheet.name) else "'" + sheet.name.replace("'", "''") + "'"
        return f"{name}!{cell}"

    @property
    def cell_count(self) -> int:
        return sum(len(s.cells) for s in self.sheets)


# -- value texts --------------------------------------------------------------------------------

_THREE_DECIMALS = re.compile(r"^-?[0-9]{1,3}\.[0-9]{3}$")
_QUOTED = re.compile(r'"[^"]*"')


def _short_float(value: float) -> str:
    if value.is_integer() and abs(value) < 1e15:
        return str(int(value))
    return format(value, ".12g")


def _positional(value: float) -> str:
    """A stored number as positional text that reads the same in every decimal style ("1.234" could be a
    thousands group, so three decimals after a short integer part get a trailing zero)."""
    if value.is_integer() and abs(value) < 1e15:
        return str(int(value))
    text = f"{value:.12f}".rstrip("0").rstrip(".")
    return text + "0" if _THREE_DECIMALS.match(text) else text


def number_texts(value: float, number_format: str | None = None) -> tuple[str, str]:
    """(text, parse text) of a stored number; a percent-formatted cell reads as its percentage ("85%")."""
    percent = "%" in _QUOTED.sub("", number_format or "")
    shown = float(value) * 100 if percent else float(value)
    if not math.isfinite(shown):
        raise ValueError("not a finite number")
    suffix = "%" if percent else ""
    shown = float(f"{shown:.12g}")  # binary noise of stored results (0.1 + 0.2)
    return _short_float(shown) + suffix, _positional(shown) + suffix


def date_text(value: datetime | date | time) -> str:
    if isinstance(value, datetime):
        if value.time() == time(0):
            return value.date().isoformat()
        return value.isoformat(sep="T", timespec="minutes")
    if isinstance(value, date):
        return value.isoformat()
    return value.isoformat(timespec="minutes")


# -- xlsx -----------------------------------------------------------------------------------------


def _check_zip(path: Path, limit: int) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            total = sum(info.file_size for info in archive.infolist())
    except (zipfile.BadZipFile, OSError) as exc:
        raise SpreadsheetError("The file could not be opened as an Excel workbook.") from exc
    if total > limit:
        raise SpreadsheetError(
            f"The workbook unpacks to more than {limit // (1024 * 1024)} MB; split it into smaller files."
        )


def _cells_of(worksheet: Any) -> Iterable[tuple[Pos, Any]]:
    """Existing cells only: a sheet with one stray cell far away is not walked cell by cell."""
    stored = getattr(worksheet, "_cells", None)
    if isinstance(stored, dict):
        return [((int(r), int(c)), cell) for (r, c), cell in stored.items()]
    return [((int(cell.row), int(cell.column)), cell) for row in worksheet.iter_rows() for cell in row]


def _xlsx_cell(pos: Pos, cell: Any) -> SheetCell | str | None:
    """A SheetCell, a reason string for an unreadable value, or None for an empty cell."""
    value = cell.value
    if value is None:
        return None
    if getattr(cell, "data_type", None) == "e":
        return f"error value {value}"
    row, col = pos
    if isinstance(value, bool):
        text = "TRUE" if value else "FALSE"
        return SheetCell(row, col, text, text, "bool", pos)
    if isinstance(value, int | float):
        try:
            text, parse = number_texts(float(value), getattr(cell, "number_format", None))
        except ValueError:
            return "not a finite number"
        return SheetCell(row, col, text, parse, "number", pos)
    if isinstance(value, datetime | date | time):
        text = date_text(value)
        return SheetCell(row, col, text, text, "date", pos)
    text = str(value).strip()
    if not text:
        return None
    return SheetCell(row, col, text, text, "text", pos)


def _read_worksheet(index: int, values_ws: Any, formulas_ws: Any | None) -> Sheet:
    formulas: dict[Pos, str] = {}
    if formulas_ws is not None:
        for pos, cell in _cells_of(formulas_ws):
            if getattr(cell, "data_type", None) == "f" and cell.value is not None:
                formulas[pos] = str(getattr(cell.value, "text", None) or cell.value)
    cells: dict[Pos, SheetCell] = {}
    pending: dict[Pos, str] = {}
    notes: dict[Pos, str] = {}
    for pos, cell in _cells_of(values_ws):
        comment = getattr(cell, "comment", None)
        if comment is not None and str(getattr(comment, "text", "") or "").strip():
            notes[pos] = " ".join(str(comment.text).split())
        read = _xlsx_cell(pos, cell)
        if isinstance(read, SheetCell):
            cells[pos] = replace(read, formula=formulas.get(pos))
        elif isinstance(read, str):
            pending[pos] = read
        elif pos in formulas:
            pending[pos] = "formula without a stored result"
    max_row = max((r for r, _ in (*cells, *pending)), default=0)
    max_col = max((c for _, c in (*cells, *pending)), default=0)
    merged: list[str] = []
    anchors: set[Pos] = set()
    for rng in list(getattr(getattr(values_ws, "merged_cells", None), "ranges", []) or []):
        top, left, bottom, right = int(rng.min_row), int(rng.min_col), int(rng.max_row), int(rng.max_col)
        merged.append(str(rng.coord))
        anchor = cells.get((top, left))
        if anchor is None:
            continue
        anchors.add((top, left))
        for r in range(top, min(bottom, max_row) + 1):
            for c in range(left, min(right, max_col) + 1):
                if (r, c) != (top, left):
                    cells[(r, c)] = replace(anchor, row=r, col=c)
    return Sheet(
        name=str(values_ws.title),
        index=index,
        cells=cells,
        pending=pending,
        notes=notes,
        merged=tuple(merged),
        merged_anchors=frozenset(anchors),
        max_row=max_row,
        max_col=max_col,
    )


def read_xlsx(path: Path) -> Workbook:
    cfg = config()
    _check_zip(path, cfg.max_uncompressed_bytes)
    import openpyxl  # imported here: only the spreadsheet stages need it

    try:
        values = openpyxl.load_workbook(path, data_only=True)  # cached formula results
        formulas = openpyxl.load_workbook(path, data_only=False)  # which cells hold formulas
    except Exception as exc:  # openpyxl raises many types for damaged or encrypted files
        logger.warning("workbook not readable: %s", type(exc).__name__)
        raise SpreadsheetError(
            "The file could not be opened as an Excel workbook. It may be damaged or encrypted."
        ) from exc
    sheets: list[Sheet] = []
    total = 0
    try:
        for index, worksheet in enumerate(values.worksheets):
            title = str(worksheet.title)
            formulas_ws = formulas[title] if title in formulas.sheetnames else None
            sheet = _read_worksheet(index, worksheet, formulas_ws)
            total += len(sheet.cells)
            if total > cfg.max_cells:
                raise SpreadsheetError(
                    f"The workbook has more than {cfg.max_cells:,} filled cells; split it into smaller files."
                )
            sheets.append(sheet)
    finally:
        values.close()
        formulas.close()
    return Workbook("xlsx", sheets)


# -- csv ------------------------------------------------------------------------------------------


def decode_text(data: bytes, encodings: Sequence[str]) -> str:
    for encoding in encodings:
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    raise SpreadsheetError("The text file's encoding could not be read.")


def sniff_delimiter(text: str, candidates: Sequence[str], lines: int) -> str:
    """The delimiter that splits the first lines into the most consistent number of fields (decimal commas
    inside semicolon-separated values split inconsistently, so they never win)."""
    sample = [line for line in text.splitlines() if line.strip()][:lines]
    best: tuple[float, int, int] | None = None
    chosen = candidates[0]
    for order, delimiter in enumerate(candidates):
        counts = [len(row) for row in csv.reader(sample, delimiter=delimiter) if any(v.strip() for v in row)]
        if not counts:
            continue
        mode, hits = Counter(counts).most_common(1)[0]
        if mode < 2:
            continue
        score = (hits / len(counts), mode, -order)
        if best is None or score > best:
            best, chosen = score, delimiter
    return chosen


def read_csv(path: Path) -> Workbook:
    cfg = config()
    text = decode_text(path.read_bytes(), cfg.csv_encodings)
    delimiter = sniff_delimiter(text, cfg.csv_delimiters, cfg.csv_sniff_lines)
    cells: dict[Pos, SheetCell] = {}
    try:
        for r, row in enumerate(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter), start=1):
            for c, value in enumerate(row, start=1):
                stripped = value.strip()
                if stripped:
                    cells[(r, c)] = SheetCell(r, c, stripped, stripped, "text", (r, c))
                    if len(cells) > cfg.max_cells:
                        raise SpreadsheetError(
                            f"The file has more than {cfg.max_cells:,} filled cells; "
                            "split it into smaller files."
                        )
    except csv.Error as exc:
        raise SpreadsheetError("The text file could not be read as a table (CSV).") from exc
    sheet = Sheet(
        name=CSV_SHEET,
        index=0,
        cells=cells,
        max_row=max((r for r, _ in cells), default=0),
        max_col=max((c for _, c in cells), default=0),
    )
    return Workbook("csv", [sheet])


def read_workbook(path: Path) -> Workbook:
    """Every sheet of an .xlsx, or the table of a .csv (by the stored file's extension)."""
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return read_xlsx(path)
    if suffix == ".csv":
        return read_csv(path)
    raise SpreadsheetError(f"Files of type {suffix or '(none)'} are not spreadsheets.")


_LAST_READ: dict[tuple[str, int, int], Workbook] = {}


def read_workbook_cached(path: Path) -> Workbook:
    """read_workbook, reusing the last file read when it is unchanged: the stages of one document read
    the same file one after the other."""
    stat = path.stat()
    key = (str(path), stat.st_size, stat.st_mtime_ns)
    found = _LAST_READ.get(key)
    if found is None:
        found = read_workbook(path)
        _LAST_READ.clear()
        _LAST_READ[key] = found
    return found


def is_spreadsheet_key(storage_key: str) -> bool:
    return Path(storage_key).suffix.lower() in EXTENSIONS


# ----------------------------------------------------------------------------------------------
# Table regions
# ----------------------------------------------------------------------------------------------


def is_number_cell(cell: SheetCell) -> bool:
    if cell.kind == "number":
        return True
    if cell.kind != "text":
        return False
    result = numbers.parse_number(cell.parse)
    return result.value is not None or result.ambiguous


def is_text_cell(cell: SheetCell) -> bool:
    return cell.kind == "text" and not is_number_cell(cell)


def _bracket_free(text: str) -> str:
    return text.strip().strip("()[]{}").strip()


def _is_unit_cell(cell: SheetCell, max_chars: int) -> bool:
    text = cell.text.strip()
    return 0 < len(text) <= max_chars and units.lookup_unit(_bracket_free(text)) is not None


@dataclass(frozen=True)
class TableGuess:
    """A table found by Python: a hint for the model and the key of a sheet layout."""

    header_row: int
    unit_row: int | None
    first_row: int
    last_row: int
    columns: tuple[int, ...]


def _distinct_sources(cells: Iterable[SheetCell]) -> int:
    return len({c.source for c in cells})


def detect_table(sheet: Sheet) -> TableGuess | None:
    """The row with the most titles that has numbers below it; a following row of units (or of more titles,
    a two-row header) is part of the header."""
    cfg = config()
    best: tuple[int, int] | None = None
    for r in sheet.content_rows()[: cfg.scan_rows]:
        cells = sheet.row_cells(r)
        texts = [c for c in cells if is_text_cell(c)]
        titles = _distinct_sources(texts)
        if titles < cfg.min_header_cells or len(texts) < cfg.min_text_share * len(cells):
            continue
        titled = {c.col for c in texts}
        below = [
            c
            for rr in range(r + 1, r + 2 + cfg.check_rows)
            for c in sheet.row_cells(rr)
            if c.col in titled and is_number_cell(c)
        ]
        if not below:
            continue
        if best is None or titles > best[0]:
            best = (titles, r)
    if best is None:
        return None
    header = best[1]
    columns = tuple(sorted({c.col for c in sheet.row_cells(header)}))
    unit_row: int | None = None
    start = header + 1
    following = [c for c in sheet.row_cells(start) if c.col in columns]
    if following and all(_is_unit_cell(c, cfg.unit_cell_max_chars) for c in following):
        unit_row, start = start, start + 1
    else:
        sub = [c for c in sheet.row_cells(start) if c.col in columns]
        if (
            sub
            and all(is_text_cell(c) for c in sub)
            and any(is_number_cell(c) for c in sheet.row_cells(start + 1) if c.col in columns)
        ):
            start += 1  # a second title row ("from", "to" under a merged "Depth")
    last = max((c.row for c in sheet.cells.values() if c.col in columns and c.row >= start), default=start)
    return TableGuess(header, unit_row, start, last, columns)


# ----------------------------------------------------------------------------------------------
# Preview: what the mapping model reads (headers and a few rows, never a whole table)
# ----------------------------------------------------------------------------------------------


@dataclass(frozen=True)
class Layout:
    """Sheets that share a table header (same row, same titles); the first is shown to the model."""

    sheet: Sheet
    table: TableGuess | None
    others: tuple[Sheet, ...] = ()

    @property
    def names(self) -> tuple[str, ...]:
        return (self.sheet.name, *(s.name for s in self.others))


@dataclass(frozen=True)
class Preview:
    text: str
    layouts: tuple[Layout, ...]
    shown: int  # layouts shown in full (the rest are listed by name only)

    @property
    def estimated_tokens(self) -> int:
        return math.ceil(len(self.text) / CHARS_PER_TOKEN)


def _fingerprint(sheet: Sheet, table: TableGuess | None) -> tuple[Any, ...] | None:
    if table is None:
        return None
    titles = tuple((c.col, _norm(c.text)) for c in sheet.row_cells(table.header_row))
    unit_cells = tuple((c.col, _norm(c.text)) for c in sheet.row_cells(table.unit_row or 0))
    return (table.header_row, table.unit_row, table.first_row, titles, unit_cells)


def layouts_of(workbook: Workbook) -> list[Layout]:
    groups: dict[tuple[Any, ...], list[tuple[Sheet, TableGuess | None]]] = {}
    order: list[tuple[Any, ...]] = []
    for sheet in workbook.sheets:
        if not sheet.cells:
            continue
        table = detect_table(sheet)
        key = _fingerprint(sheet, table) or ("sheet", sheet.index)
        if key not in groups:
            groups[key] = []
            order.append(key)
        groups[key].append((sheet, table))
    out = []
    for key in order:
        (first, table), *rest = groups[key]
        out.append(Layout(first, table, tuple(s for s, _ in rest)))
    return out


def _clip(text: str, limit: int) -> str:
    return text if len(text) <= limit else text[: max(1, limit - 1)] + "…"


def _shown(cell: SheetCell, limit: int) -> str:
    if cell.kind in ("number", "bool"):
        return cell.text
    return json.dumps(_clip(cell.text, limit), ensure_ascii=False)


def _row_line(sheet: Sheet, row: int, cfg: SheetConfig) -> str | None:
    parts: list[str] = []
    entries: list[tuple[int, str]] = [
        (c.col, f"{a1(row, c.col)}={_shown(c, cfg.max_cell_chars)}")
        for c in sheet.row_cells(row)
        if not c.merged_copy
    ]
    entries += [(col, f"{a1(row, col)}=({reason})") for (r, col), reason in sheet.pending.items() if r == row]
    for _, part in sorted(entries)[: cfg.max_columns]:
        parts.append(part)
    if not parts:
        return None
    more = len(entries) - len(parts)
    return f"R{row}: " + " | ".join(parts) + (f" | (+{more} more cells)" if more > 0 else "")


# A column is not summarised once this many of its cells are numbers and none is text.
_NUMERIC_PROBE = 20


def _category_counts(sheet: Sheet, table: TableGuess, col: int, cfg: SheetConfig) -> Counter[str] | None:
    """Counts of a short text column's values; None as soon as the column is not one (stops early)."""
    counts: Counter[str] = Counter()
    numbers_seen = 0
    for r in range(table.first_row, table.last_row + 1):
        cell = sheet.cell(r, col)
        if cell is None:
            continue
        if cell.text not in counts:
            if not is_text_cell(cell):
                numbers_seen += 1
                if numbers_seen >= _NUMERIC_PROBE and not counts:
                    return None
                continue
            if len(cell.text) > cfg.max_distinct_chars or len(counts) >= cfg.max_distinct_values:
                return None
        counts[cell.text] += 1
    return counts


def _distinct_lines(sheet: Sheet, table: TableGuess, cfg: SheetConfig) -> list[str]:
    lines = []
    for col in table.columns[: cfg.max_columns]:
        counts = _category_counts(sheet, table, col, cfg)
        if counts is None or sum(counts.values()) < 2 or len(counts) == sum(counts.values()):
            continue
        listed = ", ".join(f"{json.dumps(v, ensure_ascii=False)} x{n}" for v, n in counts.most_common())
        lines.append(f"column {column_letter(col)} values: {listed}")
    return lines


def _has_digit(cell: SheetCell) -> bool:
    return any(ch.isdigit() for ch in cell.text)


def further_tables(sheet: Sheet, table: TableGuess, cfg: SheetConfig) -> list[int]:
    """Rows below the sampled rows that look like the titles of another table (all text, numbers below),
    so the model sees every table of a sheet, not only the first."""
    found: list[int] = []
    titles = [_norm(c.text) for c in sheet.row_cells(table.header_row)]
    start = table.first_row + cfg.sample_rows
    for r in sheet.content_rows():
        if r < start or r > table.last_row:
            continue
        cells = sheet.row_cells(r)
        if sum(1 for c in cells if not _has_digit(c)) < cfg.min_header_cells:
            continue  # cheap test first: data rows hold numbers
        texts = [c for c in cells if is_text_cell(c)]
        if len(texts) < len(cells) or _distinct_sources(texts) < cfg.min_header_cells:
            continue
        if [_norm(c.text) for c in cells] == titles:
            continue  # the header printed again (page breaks)
        titled = {c.col for c in texts}
        if any(
            is_number_cell(c)
            for rr in range(r + 1, r + 2 + cfg.check_rows)
            for c in sheet.row_cells(rr)
            if c.col in titled
        ):
            found.append(r)
            if len(found) >= cfg.max_further_tables:
                break
    return found


def _section_lines(layout: Layout, cfg: SheetConfig) -> list[str]:
    sheet, table = layout.sheet, layout.table
    lines = [f"=== SHEET {json.dumps(sheet.name, ensure_ascii=False)} ==="]
    size = f"rows 1-{sheet.max_row}, columns A-{column_letter(max(sheet.max_col, 1))}"
    if sheet.merged:
        size += "; merged ranges: " + ", ".join(sheet.merged[:20])
    lines.append(size)
    if layout.others:
        names = ", ".join(json.dumps(s.name, ensure_ascii=False) for s in layout.others)
        lines.append(f"same layout (this sheet's mapping applies to them too): {names}")
    content = sheet.content_rows()
    pending_rows = sorted({r for r, _ in sheet.pending})
    all_rows = sorted(set(content) | set(pending_rows))
    shown_rows: list[int]
    if table is None:
        lines.append("no table detected")
        shown_rows = all_rows[: cfg.rows_without_table]
        hidden = len(all_rows) - len(shown_rows)
    else:
        data_rows = [r for r in all_rows if table.first_row <= r <= table.last_row]
        unit = f", unit row {table.unit_row}" if table.unit_row else ""
        lines.append(
            f"detected table: header row {table.header_row}{unit}, data rows {table.first_row}-"
            f"{table.last_row} ({len(data_rows)} rows with values)"
        )
        above = [r for r in all_rows if r < table.first_row][-(cfg.rows_above_header + 2) :]
        samples = data_rows[: cfg.sample_rows]
        further: list[int] = []
        for r in further_tables(sheet, table, cfg):
            further += [rr for rr in data_rows if r <= rr][: cfg.sample_rows + 1]
        shown_rows = sorted({*above, *samples, *further})
        hidden = len(data_rows) - len([r for r in shown_rows if r >= table.first_row])
    for r in shown_rows:
        line = _row_line(sheet, r, cfg)
        if line is not None:
            lines.append(line)
    if hidden > 0:
        lines.append(f"({hidden} more rows with values are not shown)")
    visible = set(shown_rows)
    notes = [
        f"{a1(r, c)}: {json.dumps(_clip(text, cfg.max_cell_chars), ensure_ascii=False)}"
        for (r, c), text in sorted(sheet.notes.items())
        if r in visible or (table is not None and r <= table.first_row)
    ][: cfg.max_notes]
    if notes:
        lines.append("notes: " + " | ".join(notes))
    if table is not None:
        lines.extend(_distinct_lines(sheet, table, cfg))
    return lines


def build_preview(workbook: Workbook) -> Preview:
    cfg = config()
    layouts = layouts_of(workbook)
    lines = [f"SPREADSHEET ({workbook.kind}), {len(workbook.sheets)} sheet(s):"]
    for sheet in workbook.sheets:
        size = f"rows 1-{sheet.max_row}" if sheet.cells else "empty"
        lines.append(f"- {json.dumps(sheet.name, ensure_ascii=False)}: {size}")
    shown = min(len(layouts), cfg.max_layouts)
    for layout in layouts[:shown]:
        lines.append("")
        lines.extend(_section_lines(layout, cfg))
    if shown < len(layouts):
        names = ", ".join(json.dumps(n, ensure_ascii=False) for lay in layouts[shown:] for n in lay.names)
        lines.append("")
        lines.append(f"(sheets with further layouts, not shown: {names})")
    return Preview("\n".join(lines), tuple(layouts), shown)


# ----------------------------------------------------------------------------------------------
# The mapping (strict schema for the model)
# ----------------------------------------------------------------------------------------------


class _Strict(BaseModel):
    model_config = ConfigDict(extra="forbid")


def _target_enum(schema: dict[str, Any]) -> None:
    schema["enum"] = list(config().targets)


class ColumnMap(_Strict):
    column: str = Field(description='Column letter, e.g. "C".')
    target: str = Field(description="One of the target field names.", json_schema_extra=_target_enum)
    test_name: str | None = Field(description="For test_value only: a Test Results name; else null.")
    unit: str | None = Field(description="Unit printed for this column (m, ft, %, MPa ...); null if none.")
    unit_cell: str | None = Field(description='Cell that prints the unit, e.g. "C4"; null if none.')


class HeaderCellMap(_Strict):
    label_cell: str | None = Field(description='Cell with the printed label, e.g. "A3"; null if none.')
    value_cell: str = Field(description='Cell with the value, e.g. "B3".')
    target: str = Field(description="One of the target field names.", json_schema_extra=_target_enum)
    unit: str | None = Field(description="Unit printed for this value; null if none.")
    unit_cell: str | None = Field(description="Cell that prints the unit; null if none.")


class RowTypeRule(_Strict):
    column: str = Field(description="Column letter that marks the row type.")
    values: list[str] = Field(description="Cell values exactly as printed that mark this row type.")
    row_type: RowTypeLiteral


class RegionMap(_Strict):
    sheet: str = Field(description="Sheet name exactly as listed.")
    header_row: int | None = Field(description="Row of the column titles; null for header cells only.")
    unit_row: int | None = Field(description="Row printing only units, else null.")
    first_data_row: int | None = Field(description="First data row; null: right after the header rows.")
    last_data_row: int | None = Field(description="Last data row; null: to the end of the sheet.")
    columns: list[ColumnMap]
    header_cells: list[HeaderCellMap]
    fill_down_columns: list[str] = Field(description="Columns whose blank cells repeat the value above.")
    default_row_type: RowTypeLiteral
    row_type_rules: list[RowTypeRule]
    groundwater_type: GroundwaterTypeLiteral | None
    description_language: str | None = Field(description="BCP 47 code of the descriptions, or null.")


class SheetMapping(_Strict):
    regions: list[RegionMap]
    decimal_separator: Literal[",", "."] | None = Field(
        description="Decimal separator of numbers printed as text, or null."
    )


TARGETS_PLACEHOLDER = "{{TARGETS}}"
ROW_TYPES_PLACEHOLDER = "{{ROW_TYPES}}"
TESTS_PLACEHOLDER = "{{TEST_NAMES}}"


@lru_cache
def system_prompt() -> str:
    """data/prompts/sheet_mapping.txt with the targets and row types (data/spreadsheet.yaml) and the Test
    Results names (data/test_names.yaml) filled in."""
    template = (get_settings().data_dir / "prompts" / PROMPT_FILE).read_text(encoding="utf-8")
    for placeholder in (TARGETS_PLACEHOLDER, ROW_TYPES_PLACEHOLDER, TESTS_PLACEHOLDER):
        if placeholder not in template:
            raise RuntimeError(f"{PROMPT_FILE} must contain {placeholder}")
    cfg = config()
    targets = "\n".join(f"{t.name} [{t.scope}, {t.kind}]: {t.description}" for t in cfg.targets.values())
    row_types = "\n".join(f"{name}: {text}" for name, text in cfg.row_types.items())
    tests = "\n".join(
        f"{spec.name} [{spec.unit}]: {', '.join(spec.aliases)}"
        if spec.aliases
        else f"{spec.name} [{spec.unit}]"
        for spec in units.lab_tests()
    )
    return (
        template.replace(TARGETS_PLACEHOLDER, targets)
        .replace(ROW_TYPES_PLACEHOLDER, row_types)
        .replace(TESTS_PLACEHOLDER, tests)
    )


def output_tokens_for(preview: Preview) -> int:
    cfg = config()
    wanted = cfg.tokens_floor + cfg.tokens_per_layout * max(1, preview.shown)
    return min(cfg.tokens_cap, wanted)


class SheetMapper(Protocol):
    async def map_sheets(
        self, preview: Preview, *, document_id: str
    ) -> tuple[SheetMapping, str | None, float]: ...


class OpenRouterSheetMapper:
    """The extractor model on OpenRouter with the strict SheetMapping schema (reasoning off by the client's
    default; the budget is checked by the client before the call)."""

    def __init__(self) -> None:
        from app.providers.llm_openrouter import OpenRouterClient

        self.client = OpenRouterClient()

    async def map_sheets(
        self, preview: Preview, *, document_id: str
    ) -> tuple[SheetMapping, str | None, float]:
        model = get_settings().extractor_model
        max_tokens = output_tokens_for(preview)
        result = await self.client.structured(
            model=model,
            output=SheetMapping,
            system=system_prompt(),
            content=[text_part(preview.text)],
            purpose="sheet_mapping",
            document_id=document_id,
            max_tokens=max_tokens,
            estimate_usd=token_cost(
                model,
                preview.estimated_tokens + math.ceil(len(system_prompt()) / CHARS_PER_TOKEN),
                max_tokens,
            ),
        )
        return result.data, result.model, result.cost_usd


def mapping_cache_key(sha256: str, model: str, preview: Preview) -> str:
    """sha256 over the file, the model, the prompt and schema versions and everything the model reads."""
    content = hashlib.sha256()
    for piece in (system_prompt(), preview.text):
        content.update(piece.encode("utf-8"))
        content.update(b"\x00")
    parts = [sha256, MAPPING_GROUP, model, PROMPT_VERSION, MAPPING_SCHEMA_VERSION, content.hexdigest()]
    return hashlib.sha256("\x1f".join(parts).encode("utf-8")).hexdigest()


@dataclass(frozen=True)
class MappingRun:
    id: str
    mapping: SheetMapping
    model: str | None
    cost_usd: float
    reused: bool


async def _stored_mapping(key: str) -> MappingRun | None:
    async with get_sessionmaker()() as session:
        run = await session.scalar(select(ExtractionRun).where(ExtractionRun.cache_key == key))
        if run is None or run.status != RUN_DONE or run.output is None:
            return None
        try:
            mapping = SheetMapping.model_validate(run.output)
        except ValueError:
            logger.warning("stored sheet mapping %s does not validate; mapping again", run.id)
            return None
        return MappingRun(run.id, mapping, run.model, float(run.cost_usd or 0.0), True)


async def ensure_mapping(
    document_id: str, sha256: str, preview: Preview, *, mapper: SheetMapper | None = None
) -> MappingRun:
    """The stored mapping for this file and preview, or one model call whose answer is stored."""
    model = get_settings().extractor_model
    key = mapping_cache_key(sha256, model, preview)
    stored = await _stored_mapping(key)
    if stored is not None:
        return stored
    async with get_sessionmaker()() as session:
        run = await session.scalar(select(ExtractionRun).where(ExtractionRun.cache_key == key))
        if run is None:
            run = ExtractionRun(
                id=new_id("ext"),
                document_id=document_id,
                group_id=MAPPING_GROUP,
                cache_key=key,
                model=model,
                prompt_version=PROMPT_VERSION,
                schema_version=MAPPING_SCHEMA_VERSION,
                input_tokens=preview.estimated_tokens,
                status="running",
            )
            session.add(run)
        else:
            run.status, run.error = "running", None
        await session.commit()
        run_id = run.id
    try:
        active = mapper if mapper is not None else OpenRouterSheetMapper()
        mapping, returned_model, cost = await active.map_sheets(preview, document_id=document_id)
    except ProviderError as exc:
        async with get_sessionmaker()() as session:
            failed = await session.get(ExtractionRun, run_id)
            if failed is not None:
                failed.status, failed.error = "failed", str(exc)
                await session.commit()
        raise
    async with get_sessionmaker()() as session:
        done = await session.get(ExtractionRun, run_id)
        if done is not None:
            done.status, done.error = RUN_DONE, None
            done.output = mapping.model_dump(mode="json")
            done.model = returned_model or model
            done.cost_usd = cost
            await session.commit()
    return MappingRun(run_id, mapping, returned_model or model, cost, False)


# ----------------------------------------------------------------------------------------------
# Validating the mapping against the workbook
# ----------------------------------------------------------------------------------------------


@dataclass(frozen=True)
class UnitUse:
    key: str  # canonical unit (data/units.yaml)
    text: str  # as printed
    pos: Pos | None  # where it is printed


@dataclass(frozen=True)
class ColumnRule:
    col: int
    target: str
    kind: TargetKind
    test_name: str | None
    unit: UnitUse | None
    title: str | None  # the column title as printed
    title_pos: Pos | None


@dataclass(frozen=True)
class HeaderRule:
    target: str
    kind: TargetKind
    value_pos: Pos
    label_pos: Pos | None
    label_text: str | None
    unit: UnitUse | None


@dataclass(frozen=True)
class RegionRule:
    sheets: tuple[str, ...]  # the sheet the model saw, then the sheets that share its layout
    header_row: int | None
    unit_row: int | None
    first_row: int | None
    last_row: int | None
    columns: tuple[ColumnRule, ...]
    header_cells: tuple[HeaderRule, ...]
    fill_down: frozenset[int]
    row_rules: tuple[tuple[int, frozenset[str], str], ...]
    default_row_type: str
    groundwater_type: str | None
    language: str | None


@dataclass
class Plan:
    regions: list[RegionRule]
    decimal_separator: str | None
    warnings: list[str] = field(default_factory=list)


def _test_dimension(test_name: str) -> str | None:
    spec = next((s for s in units.lab_tests() if s.name == test_name), None)
    unit = units.lookup_unit(spec.unit) if spec is not None else None
    return unit.dimension if unit is not None else None


def unit_in(text: str | None, dimension: str | None) -> tuple[str, str] | None:
    """(canonical unit, printed text) a title, unit cell or note names for `dimension` (any when None):
    bracketed text first, then the whole text, then the last word that names such a unit."""
    if text is None or not text.strip():
        return None
    if dimension == "length":
        key = length_unit_in(text)
        if key is None:
            return None
        return key, text.strip()
    candidates = [m.group(1) for m in re.finditer(r"[(\[{]([^)\]}]*)[)\]}]", text)]
    for candidate in [*candidates, text]:
        resolved = units.resolve_unit(candidate)
        if resolved is not None and (dimension is None or resolved[0].dimension == dimension):
            return resolved[0].key, resolved[1]
    found = None
    for word in text.split():
        resolved = units.resolve_unit(word.strip(".,;:"))
        if resolved is not None and (dimension is None or resolved[0].dimension == dimension):
            found = (resolved[0].key, resolved[1])
    return found


def _title_cells(sheet: Sheet, col: int, header_row: int | None, unit_row: int | None) -> list[SheetCell]:
    """Cells that title a column: the header cell, the unit-row cell and a merged title just above."""
    found: list[SheetCell] = []
    for row in (header_row, unit_row):
        cell = sheet.cell(row, col) if row is not None else None
        if cell is not None:
            found.append(cell)
    if header_row is not None and header_row > 1:
        above = sheet.cell(header_row - 1, col)
        if above is not None and (above.merged_copy or above.source in sheet.merged_anchors):
            found.append(above)
    return found


def _source_text(sheet: Sheet, pos: Pos) -> tuple[str | None, str | None]:
    cell = sheet.cell(*pos)
    return (cell.text if cell is not None else None), sheet.notes.get(pos)


class _Validator:
    def __init__(self, workbook: Workbook, preview: Preview) -> None:
        self.workbook = workbook
        self.cfg = config()
        self.layouts = {layout.sheet.name: layout for layout in preview.layouts}
        self.warnings: list[str] = []

    def warn(self, message: str) -> None:
        self.warnings.append(message)

    def unit_for(
        self,
        sheet: Sheet,
        where: str,
        kind: TargetKind,
        test_name: str | None,
        printed_cells: Sequence[SheetCell],
        model_unit: str | None,
        model_cell: str | None,
    ) -> UnitUse | None:
        dimension = _test_dimension(test_name) if kind == "test" and test_name else _KIND_DIMENSION.get(kind)
        if kind not in ("test",) and dimension is None:
            return None
        for cell in printed_cells:
            hit = unit_in(cell.text, dimension)
            if hit is not None:
                if (
                    model_unit
                    and (resolved := units.lookup_unit(model_unit)) is not None
                    and resolved.key != hit[0]
                ):
                    self.warn(f"{where}: the model named unit '{model_unit}'; the printed '{hit[1]}' is used")
                return UnitUse(hit[0], hit[1], cell.source)
        if not model_unit:
            return None
        pos = parse_ref(model_cell)
        if pos is not None:
            text, note = _source_text(sheet, pos)
            wanted = units.lookup_unit(model_unit)
            for candidate in (text, note):
                hit = unit_in(candidate, dimension)
                if hit is not None and wanted is not None and hit[0] == wanted.key:
                    at = sheet.cell(*pos)
                    return UnitUse(hit[0], hit[1], at.source if at is not None else pos)
        self.warn(f"{where}: unit '{model_unit}' is not printed where the model says; no unit used")
        return None

    def target(self, name: str, where: str, *, header: bool) -> TargetSpec | None:
        spec = self.cfg.targets.get(name)
        if spec is None:
            self.warn(f"{where}: '{name}' is not a target field; ignored")
            return None
        if header and spec.scope == "row":
            self.warn(f"{where}: '{name}' is a per-row field and cannot come from a header cell; ignored")
            return None
        return spec

    def region(self, region: RegionMap) -> RegionRule | None:
        sheet = self.workbook.sheet(region.sheet)
        if sheet is None:
            self.warn(f"region on sheet '{region.sheet}': no such sheet; ignored")
            return None
        layout = self.layouts.get(sheet.name)
        sheets = layout.names if layout is not None else (sheet.name,)
        where = f"sheet '{sheet.name}'"
        bound = max(sheet.max_row, 1)
        header_row = (
            region.header_row if region.header_row is not None and 1 <= region.header_row <= bound else None
        )
        if region.header_row is not None and header_row is None:
            self.warn(f"{where}: header row {region.header_row} is outside the sheet; table ignored")
        unit_row = region.unit_row if region.unit_row is not None and 1 <= region.unit_row <= bound else None
        if unit_row == header_row:
            unit_row = None
        first_row: int | None = None
        last_row: int | None = None
        if header_row is not None:
            first_row = max(header_row, unit_row or 0) + 1
            if region.first_data_row is not None and region.first_data_row > header_row:
                first_row = region.first_data_row
            if region.last_data_row is not None and region.last_data_row >= first_row:
                last_row = region.last_data_row
        columns = self.columns(sheet, region, header_row, unit_row, where)
        header_cells = self.header_cells(sheet, region, where)
        has_id = any(c.target == "borehole_id" for c in columns) or any(
            h.target == "borehole_id" for h in header_cells
        )
        if not has_id:
            self.warn(f"{where}: no borehole id column or cell; region ignored")
            return None
        fill_down = frozenset(
            c for c in (column_index(letter) for letter in region.fill_down_columns) if c is not None
        )
        rules: list[tuple[int, frozenset[str], str]] = []
        for rule in region.row_type_rules:
            col = column_index(rule.column)
            values = frozenset(_norm(v) for v in rule.values if v.strip())
            if col is None or col > sheet.max_col or not values:
                self.warn(f"{where}: row-type rule on column '{rule.column}' is not usable; ignored")
                continue
            rules.append((col, values, rule.row_type))
        language = (region.description_language or "").strip() or None
        return RegionRule(
            sheets=sheets,
            header_row=header_row,
            unit_row=unit_row,
            first_row=first_row,
            last_row=last_row,
            columns=tuple(columns),
            header_cells=tuple(header_cells),
            fill_down=fill_down,
            row_rules=tuple(rules),
            default_row_type=region.default_row_type,
            groundwater_type=region.groundwater_type,
            language=language,
        )

    def columns(
        self, sheet: Sheet, region: RegionMap, header_row: int | None, unit_row: int | None, where: str
    ) -> list[ColumnRule]:
        out: list[ColumnRule] = []
        if header_row is None:
            if region.columns:
                self.warn(f"{where}: columns without a header row; ignored")
            return out
        seen: set[str] = set()
        used_columns: set[int] = set()
        for item in region.columns:
            label = f"{where} column {item.column}"
            col = column_index(item.column)
            if col is None or col > sheet.max_col:
                self.warn(f"{label}: no such column; ignored")
                continue
            spec = self.target(item.target, label, header=False)
            if spec is None:
                continue
            titles = _title_cells(sheet, col, header_row, unit_row)
            if not titles:
                self.warn(f"{label}: no title is printed for this column; ignored")
                continue
            if col in used_columns:
                self.warn(f"{label}: already mapped to another target; ignored")
                continue
            if spec.name in seen and spec.name not in MULTI_COLUMN_TARGETS:
                self.warn(f"{label}: '{spec.name}' is already mapped to another column; ignored")
                continue
            test_name: str | None = None
            if spec.kind == "test":
                test_name = units.match_lab_test(item.test_name)
                if test_name is None:
                    self.warn(f"{label}: '{item.test_name}' is not a Test Results name; ignored")
                    continue
            unit = self.unit_for(sheet, label, spec.kind, test_name, titles, item.unit, item.unit_cell)
            header_cell = sheet.cell(header_row, col)
            out.append(
                ColumnRule(
                    col=col,
                    target=spec.name,
                    kind=spec.kind,
                    test_name=test_name,
                    unit=unit,
                    title=header_cell.text if header_cell is not None else titles[0].text,
                    title_pos=(header_cell or titles[0]).source,
                )
            )
            seen.add(spec.name)
            used_columns.add(col)
        return out

    def header_cells(self, sheet: Sheet, region: RegionMap, where: str) -> list[HeaderRule]:
        out: list[HeaderRule] = []
        seen: set[str] = set()
        for item in region.header_cells:
            label = f"{where} cell {item.value_cell}"
            pos = parse_ref(item.value_cell)
            if pos is None or pos[0] > sheet.max_row or pos[1] > sheet.max_col:
                self.warn(f"{label}: no such cell; ignored")
                continue
            spec = self.target(item.target, label, header=True)
            if spec is None:
                continue
            if spec.name in seen:
                self.warn(f"{label}: '{spec.name}' already comes from another cell; ignored")
                continue
            label_pos = parse_ref(item.label_cell) if item.label_cell else None
            label_cell = sheet.cell(*label_pos) if label_pos is not None else None
            if item.label_cell and label_cell is None:
                self.warn(f"{label}: label cell {item.label_cell} is empty; the value is read without it")
            printed = [c for c in (label_cell,) if c is not None]
            unit = self.unit_for(sheet, label, spec.kind, None, printed, item.unit, item.unit_cell)
            out.append(
                HeaderRule(
                    target=spec.name,
                    kind=spec.kind,
                    value_pos=pos,
                    label_pos=label_cell.source if label_cell is not None else None,
                    label_text=label_cell.text if label_cell is not None else None,
                    unit=unit,
                )
            )
            seen.add(spec.name)
        return out


def _clip_overlaps(regions: list[RegionRule], warnings: list[str]) -> list[RegionRule]:
    """Tables on one sheet never share rows: an earlier table ends above the next one's header."""
    out = list(regions)
    by_sheet: dict[str, list[int]] = {}
    for i, region in enumerate(out):
        if region.header_row is not None:
            by_sheet.setdefault(region.sheets[0], []).append(i)
    for indexes in by_sheet.values():
        indexes.sort(key=lambda i: out[i].header_row or 0)
        for current, following in itertools.pairwise(indexes):
            upper, lower = out[current], out[following]
            assert lower.header_row is not None
            end = lower.header_row - 1
            if upper.last_row is None or upper.last_row > end:
                if upper.first_row is not None and end < upper.first_row:
                    warnings.append(f"sheet '{upper.sheets[0]}': two tables share a header row; one ignored")
                    out[following] = replace(lower, columns=(), header_cells=())
                    continue
                out[current] = replace(upper, last_row=end)
    return [r for r in out if r.columns or r.header_cells]


def validate_mapping(mapping: SheetMapping, workbook: Workbook, preview: Preview) -> Plan:
    """The usable part of the model's mapping; everything dropped is named in `warnings`."""
    validator = _Validator(workbook, preview)
    regions = [r for r in (validator.region(region) for region in mapping.regions) if r is not None]
    # A sheet the model mapped itself is not also read through another sheet's layout.
    explicit = {r.sheets[0] for r in regions}
    regions = [
        replace(r, sheets=(r.sheets[0], *(s for s in r.sheets[1:] if s not in explicit))) for r in regions
    ]
    regions = _clip_overlaps(regions, validator.warnings)
    return Plan(regions, mapping.decimal_separator, validator.warnings)


# ----------------------------------------------------------------------------------------------
# Applying the mapping to every row
# ----------------------------------------------------------------------------------------------


@dataclass(frozen=True)
class CellValue:
    ref: str  # provenance cell
    text: str  # the stored value as text
    parse: str  # what the normalisers read (the column's printed unit appended when the value has none)
    order: tuple[int, int, int]  # sheet index, row, column of the printed cell
    title: bool = False  # a column title, unit or label cell: cited, never the value's own cell


@dataclass
class SheetRow:
    sheet: str
    row: int
    row_type: str
    values: dict[str, list[CellValue]]
    tests: list[tuple[str, CellValue, CellValue | None]]  # Test Results name, value, printed unit
    groundwater_type: str | None
    language: str | None

    def first(self, target: str) -> CellValue | None:
        found = self.values.get(target)
        return found[0] if found else None


@dataclass
class BoreholeRows:
    key: str  # canonical borehole id
    label: CellValue
    header: dict[str, CellValue] = field(default_factory=dict)
    axis_labels: dict[str, CellValue] = field(default_factory=dict)  # titles or labels of x and y
    rows: list[SheetRow] = field(default_factory=list)
    conflicts: list[str] = field(default_factory=list)
    groundwater_type: str | None = None  # of a groundwater value printed in a header block


_NUMERIC_KINDS = frozenset({"length", "percent", "angle", "test"})


def _glue(cell: SheetCell, kind: TargetKind, unit: UnitUse | None) -> str:
    """What the normalisers read for a cell of a `kind` column: the stored text for texts, dates and
    coordinates (kept exactly as stored); for measured values the parser form of a stored number, with the
    column's printed unit appended when the value prints none itself."""
    if kind not in _NUMERIC_KINDS:
        return cell.text
    if unit is None or kind == "test":
        return cell.parse
    parsed = numbers.parse_number(cell.parse)
    if parsed.unit is not None or (parsed.value is None and not parsed.ambiguous):
        return cell.parse
    return f"{cell.parse} {unit.key}"


class _Applier:
    def __init__(self, workbook: Workbook, plan: Plan) -> None:
        self.workbook = workbook
        self.plan = plan
        self.cfg = config()
        self.boreholes: dict[str, BoreholeRows] = {}
        self.warnings: list[str] = []
        self.unread: dict[tuple[str, str, str], list[str]] = {}
        self.not_ids: list[str] = []

    # -- cells ----------------------------------------------------------------------------------

    def value(
        self, sheet: Sheet, cell: SheetCell, parse: str | None = None, *, title: bool = False
    ) -> CellValue:
        return CellValue(
            ref=self.workbook.ref(sheet, cell.source),
            text=cell.text,
            parse=parse if parse is not None else cell.text,
            order=(sheet.index, *cell.source),
            title=title,
        )

    def unit_value(self, sheet: Sheet, unit: UnitUse | None) -> CellValue | None:
        if unit is None or unit.pos is None:
            return None
        cell = sheet.cell(*unit.pos)
        if cell is None:
            return None
        return CellValue(
            self.workbook.ref(sheet, cell.source), unit.text, unit.text, (sheet.index, *cell.source), True
        )

    def pending(self, sheet: Sheet, pos: Pos) -> None:
        reason = sheet.pending.get(pos)
        if reason is not None:
            key = (sheet.name, column_letter(pos[1]), reason)
            self.unread.setdefault(key, []).append(self.workbook.ref(sheet, pos))

    def summary(self) -> list[str]:
        """One warning per column and reason for cells whose value is not stored (never one per row)."""
        out = []
        for (name, letter, reason), refs in self.unread.items():
            shown = ", ".join(refs[:5]) + (" ..." if len(refs) > 5 else "")
            out.append(
                f"sheet '{name}' column {letter}: {len(refs)} cell(s) with {reason} not read ({shown})"
            )
        if self.not_ids:
            shown = ", ".join(self.not_ids[:5]) + (" ..." if len(self.not_ids) > 5 else "")
            out.append(
                f"{len(self.not_ids)} borehole id cell(s) hold no usable id; their rows skipped ({shown})"
            )
        return out

    # -- boreholes ------------------------------------------------------------------------------

    def borehole(self, label: CellValue) -> BoreholeRows | None:
        key = canonical_label(label.text)
        if not key:
            self.not_ids.append(label.ref)
            return None
        found = self.boreholes.get(key)
        if found is None:
            found = BoreholeRows(key=key, label=label)
            self.boreholes[key] = found
        return found

    def set_header(self, bh: BoreholeRows, target: str, value: CellValue) -> None:
        existing = bh.header.get(target)
        if existing is None:
            bh.header[target] = value
            return
        if _norm(existing.text) != _norm(value.text) and target not in {
            c.split(":")[0] for c in bh.conflicts
        }:
            bh.conflicts.append(f"{target}: {existing.ref} '{existing.text}' vs {value.ref} '{value.text}'")
            self.warnings.append(
                f"borehole {bh.label.text}: {target} printed as '{existing.text}' ({existing.ref}) and "
                f"'{value.text}' ({value.ref}); the first is used"
            )

    # -- header cells ---------------------------------------------------------------------------

    def locate(self, sheet: Sheet, rule: HeaderRule, region: RegionRule) -> Pos | None:
        """The value cell of a header rule on `sheet`: the same place, or the same offset from its label
        when the label sits elsewhere on a sheet of the same layout."""
        if rule.label_pos is None or rule.label_text is None:
            return rule.value_pos
        here = sheet.cell(*rule.label_pos)
        if here is not None and _norm(here.text) == _norm(rule.label_text):
            return rule.value_pos
        limit = region.header_row if region.header_row is not None else sheet.max_row + 1
        for r in sheet.content_rows():
            if r >= limit:
                break
            for cell in sheet.row_cells(r):
                if not cell.merged_copy and _norm(cell.text) == _norm(rule.label_text):
                    dr, dc = rule.value_pos[0] - rule.label_pos[0], rule.value_pos[1] - rule.label_pos[1]
                    target = (cell.row + dr, cell.col + dc)
                    return target if target[0] >= 1 and target[1] >= 1 else None
        self.warnings.append(
            f"sheet '{sheet.name}': label '{rule.label_text}' not found; {rule.target} not read"
        )
        return None

    def header_values(
        self, sheet: Sheet, region: RegionRule
    ) -> tuple[dict[str, CellValue], dict[str, CellValue]]:
        values: dict[str, CellValue] = {}
        labels: dict[str, CellValue] = {}
        for rule in region.header_cells:
            pos = self.locate(sheet, rule, region)
            if pos is None:
                continue
            cell = sheet.cell(*pos)
            if cell is None:
                self.pending(sheet, pos)
                continue
            values[rule.target] = self.value(sheet, cell, _glue(cell, rule.kind, rule.unit))
            if rule.target in ("x", "y") and rule.label_pos is not None:
                delta = (pos[0] - rule.value_pos[0], pos[1] - rule.value_pos[1])
                label_cell = sheet.cell(rule.label_pos[0] + delta[0], rule.label_pos[1] + delta[1])
                if label_cell is not None:
                    labels[rule.target] = self.value(sheet, label_cell, title=True)
        return values, labels

    # -- rows -----------------------------------------------------------------------------------

    def row_type(self, sheet: Sheet, row: int, region: RegionRule) -> str:
        for col, values, row_type in region.row_rules:
            cell = sheet.cell(row, col)
            if cell is not None and _norm(cell.text) in values:
                return row_type
        return region.default_row_type

    def header_repeat(self, sheet: Sheet, row: int, region: RegionRule) -> bool:
        titled = [c for c in region.columns if c.title]
        if not titled:
            return False
        same = sum(
            1
            for c in titled
            if (cell := sheet.cell(row, c.col)) is not None and _norm(cell.text) == _norm(c.title)
        )
        return same >= max(2, math.ceil(self.cfg.header_repeat_share * len(titled)))

    def apply_region(self, region: RegionRule) -> None:
        for name in region.sheets:
            sheet = self.workbook.sheet(name)
            if sheet is None:
                continue
            header, labels = self.header_values(sheet, region)
            sheet_label = header.pop("borehole_id", None)
            sheet_bh = self.borehole(sheet_label) if sheet_label is not None else None
            touched: list[BoreholeRows] = [sheet_bh] if sheet_bh is not None else []
            if region.header_row is not None and region.columns:
                touched += self.apply_table(sheet, region, sheet_bh)
            for bh in {id(b): b for b in touched}.values():
                for target, value in header.items():
                    self.set_header(bh, target, value)
                for axis, value in labels.items():
                    bh.axis_labels.setdefault(axis, value)
                if "groundwater_depth" in header and bh.groundwater_type is None:
                    bh.groundwater_type = region.groundwater_type

    def apply_table(
        self, sheet: Sheet, region: RegionRule, sheet_bh: BoreholeRows | None
    ) -> list[BoreholeRows]:
        assert region.first_row is not None
        last = region.last_row if region.last_row is not None else sheet.max_row
        id_rule = next((c for c in region.columns if c.target == "borehole_id"), None)
        rule_columns = {col for col, _, _ in region.row_rules}
        carried: dict[int, SheetCell] = {}
        touched: list[BoreholeRows] = []
        unlabelled: list[int] = []
        rows = [r for r in sheet.content_rows() if region.first_row <= r <= last]
        pending_rows = sorted({r for r, _ in sheet.pending if region.first_row <= r <= last})
        id_col = id_rule.col if id_rule is not None else None
        for row in sorted(set(rows) | set(pending_rows)):
            cells: dict[int, SheetCell] = {}
            own: set[int] = set()
            for rule in region.columns:
                cell = sheet.cell(row, rule.col)
                if cell is not None:
                    own.add(rule.col)
                    if rule.col in region.fill_down:
                        carried[rule.col] = cell
                elif rule.col in region.fill_down:
                    cell = carried.get(rule.col)
                if cell is not None:
                    cells[rule.col] = cell
            row_type = self.row_type(sheet, row, region)
            if row_type == "ignore" or self.header_repeat(sheet, row, region):
                continue
            for rule in region.columns:
                if rule.col not in cells:
                    self.pending(sheet, (row, rule.col))
            if not any(col != id_col and col not in rule_columns for col in own):
                continue  # nothing but an id or a row-type mark (or values carried down from above)
            label_cell = cells.get(id_col) if id_col is not None else None
            bh = self.borehole(self.value(sheet, label_cell)) if label_cell is not None else sheet_bh
            if bh is None:
                unlabelled.append(row)
                continue
            touched.append(bh)
            self.add_row(bh, sheet, row, row_type, region, cells)
        if unlabelled:
            shown = ", ".join(str(r) for r in unlabelled[:10])
            self.warnings.append(
                f"sheet '{sheet.name}': {len(unlabelled)} row(s) with values but no borehole id skipped "
                f"(rows {shown}{' ...' if len(unlabelled) > 10 else ''})"
            )
        return touched

    def add_row(
        self,
        bh: BoreholeRows,
        sheet: Sheet,
        row: int,
        row_type: str,
        region: RegionRule,
        cells: Mapping[int, SheetCell],
    ) -> None:
        item = SheetRow(sheet.name, row, row_type, {}, [], region.groundwater_type, region.language)
        for rule in region.columns:
            cell = cells.get(rule.col)
            if cell is None or rule.target == "borehole_id":
                continue
            spec = self.cfg.targets[rule.target]
            value = self.value(sheet, cell, _glue(cell, rule.kind, rule.unit))
            if spec.kind == "test" and rule.test_name is not None:
                glued = numbers.parse_number(cell.parse).unit is not None
                unit = None if glued else self.unit_value(sheet, rule.unit)
                item.tests.append((rule.test_name, value, unit))
            elif spec.scope == "borehole":
                self.set_header(bh, rule.target, value)
                if rule.target in ("x", "y") and rule.title_pos is not None:
                    title = sheet.cell(*rule.title_pos)
                    if title is not None:
                        bh.axis_labels.setdefault(rule.target, self.value(sheet, title, title=True))
            else:
                item.values.setdefault(rule.target, []).append(value)
        if item.values or item.tests:
            bh.rows.append(item)


def apply_plan(workbook: Workbook, plan: Plan) -> tuple[list[BoreholeRows], list[str]]:
    """Every row of every mapped region, grouped by borehole id in order of first appearance."""
    applier = _Applier(workbook, plan)
    for region in plan.regions:
        applier.apply_region(region)
    return list(applier.boreholes.values()), [*applier.warnings, *applier.summary()]


# ----------------------------------------------------------------------------------------------
# Records: rows -> ExtractionResult -> app.pipeline.finalize.build_record -> cell provenance
# ----------------------------------------------------------------------------------------------

# Synthetic "pages" of title cells sort after every value cell, so an entry that cites a value and its
# column's unit points at the value's cell.
_TITLE_PAGE_OFFSET = 1_000_000


class _Cells:
    """Cells cited by one borehole's ExtractionResult, as OCR-like blocks with one synthetic page each.

    The page numbers never leave this module: provenance entries are turned back into cells."""

    def __init__(self) -> None:
        self.blocks: dict[str, BlockRef] = {}
        self.ref_of_page: dict[int, str] = {}
        self.text_of: dict[str, str] = {}  # parse text -> stored text
        self._values = 0
        self._titles = 0

    def cite(self, value: CellValue) -> str:
        if value.ref not in self.blocks:
            if value.title:
                self._titles += 1
                page = _TITLE_PAGE_OFFSET + self._titles
            else:
                self._values += 1
                page = self._values
            self.blocks[value.ref] = BlockRef(
                id=value.ref,
                page=page,
                text=value.text,
                box=None,
                confidence=None,
                page_text_quality=None,
                page_handwriting=None,
            )
            self.ref_of_page[page] = value.ref
        if value.parse != value.text:
            self.text_of.setdefault(value.parse, value.text)
        return value.ref

    def raw(self, value: CellValue | None) -> RawValue:
        if value is None:
            return RawValue.absent()
        return RawValue(raw=value.parse, block_ids=[self.cite(value)], state="ok")

    def joined(self, values: Sequence[CellValue] | None, separator: str) -> RawValue:
        if not values:
            return RawValue.absent()
        if len(values) == 1:
            return self.raw(values[0])
        return RawValue(
            raw=separator.join(v.parse for v in values), block_ids=[self.cite(v) for v in values], state="ok"
        )


@dataclass
class SheetDraft:
    """One borehole of a spreadsheet, ready for create_borehole."""

    key: str
    label: str
    body: dict[str, Any] | None
    provenance: dict[str, ProvenanceIn]
    crs_status: str
    crs_code: str | None
    guessed: list[tuple[str, str, float]]  # (cell, text, metres) of depths stored as metres without a unit
    warnings: list[str]


class _Result:
    """Builds the ExtractionResult of one borehole and remembers what the post-processing needs."""

    def __init__(self, bh: BoreholeRows) -> None:
        self.bh = bh
        self.cells = _Cells()
        self.join = config().text_join
        # Depth cells of layer rows used to attach the row's SPT to its layer (the SPT has no depth).
        self.layer_spt_depths: set[str] = set()
        self.din_cells: dict[str, CellValue] = {}  # layer top cell -> printed DIN code cell

    def build(self) -> ExtractionResult:
        c, h = self.cells, self.bh.header
        labels = RawValue.absent()
        x_label, y_label = self.bh.axis_labels.get("x"), self.bh.axis_labels.get("y")
        if x_label is not None and y_label is not None:
            labels = RawValue(
                raw=f"{x_label.text}, {y_label.text}",
                block_ids=[c.cite(x_label), c.cite(y_label)],
                state="ok",
            )
        header = ExtractedHeader(
            borehole_label=c.raw(self.bh.label),
            borehole_name=c.raw(h.get("borehole_name")),
            project=c.raw(h.get("project_name")),
            project_id=c.raw(h.get("project_id")),
            client=c.raw(h.get("client")),
            location=c.raw(h.get("location_name")),
            contractor=c.raw(h.get("drilling_contractor")),
            method=c.raw(h.get("drilling_method")),
            start_date=c.raw(h.get("start_date")),
            end_date=c.raw(h.get("end_date")),
            total_depth=c.raw(h.get("total_depth")),
            borehole_length=c.raw(h.get("borehole_length")),
            depth_unit=RawValue.absent(),  # each depth carries its column's printed unit
            ground_level=c.raw(h.get("ground_level")),
            elevation_datum=c.raw(h.get("elevation_datum")),
            coordinates=RawCoordinates(
                x_raw=c.raw(h.get("x")),
                y_raw=c.raw(h.get("y")),
                labels_raw=labels,
                system_raw=c.raw(h.get("coordinate_system")),
            ),
            inclination=c.raw(h.get("inclination")),
            depth_origin_note=RawValue.absent(),
            notes=c.raw(h.get("notes")),
        )
        layers: list[ExtractedLayer] = []
        spt: list[ExtractedSpt] = []
        runs: list[ExtractedCoreRun] = []
        water: list[ExtractedGroundwater] = []
        samples: list[ExtractedSample] = []
        tests: list[ExtractedTest] = []
        for row in self.bh.rows:
            self.layer(row, layers)
            self.spt(row, spt)
            self.core_run(row, runs)
            self.groundwater(row, water)
            self.sample(row, samples)
            self.tests(row, tests)
        if "groundwater_depth" in h:  # printed once in a header block
            water.append(
                ExtractedGroundwater(
                    type=self._water_type(None),
                    depth_raw=c.raw(h["groundwater_depth"]),
                    date_raw=c.raw(h.get("groundwater_date")),
                )
            )
        return ExtractionResult(
            header=header,
            layers=layers,
            spt=spt,
            core_runs=runs,
            groundwater=water,
            samples=samples,
            tests=tests,
        )

    def _water_type(self, row: SheetRow | None) -> GroundwaterTypeLiteral:
        value = row.groundwater_type if row is not None else self.bh.groundwater_type
        return cast(GroundwaterTypeLiteral, value) if value in get_args(GroundwaterTypeLiteral) else "other"

    def layer(self, row: SheetRow, out: list[ExtractedLayer]) -> None:
        if row.row_type != "layer":
            return
        c = self.cells
        top, base = row.first("depth_from"), row.first("depth_to")
        if top is None and base is None and "description" not in row.values:
            return
        code = row.first("din4023_code")
        if code is not None and top is not None:
            self.din_cells[top.ref] = code
        out.append(
            ExtractedLayer(
                from_raw=c.raw(top),
                to_raw=c.raw(base),
                description_raw=c.joined(row.values.get("description"), self.join),
                description_language=row.language,
                english_lithology=None,
                din4023_code=code.text.strip() if code is not None else None,
                din4023_notation=None,
                uscs_raw=c.raw(row.first("uscs")),
                colour_raw=c.raw(row.first("colour")),
                consistency_raw=c.raw(row.first("consistency")),
                moisture_raw=c.raw(row.first("moisture")),
                weathering_raw=c.raw(row.first("weathering")),
                stratigraphic_name_raw=c.raw(row.first("stratigraphic_name")),
                engineering_soil_unit_raw=c.raw(row.first("engineering_soil_unit")),
                geologic_age_raw=c.raw(row.first("geologic_age")),
            )
        )

    def spt(self, row: SheetRow, out: list[ExtractedSpt]) -> None:
        blows, n_value = row.first("spt_blows"), row.first("spt_n")
        if blows is None and n_value is None:
            return
        depth = row.first("spt_depth")
        if depth is None and row.row_type == "layer":
            depth = row.first("depth_from")
            if depth is not None:
                self.layer_spt_depths.add(depth.ref)  # attaches the SPT to its row's layer; not its depth
        elif depth is None and row.row_type != "borehole":
            depth = row.first("depth_from")
        c = self.cells
        out.append(
            ExtractedSpt(
                depth_raw=c.raw(depth),
                blows_raw=c.raw(blows),
                n_raw=c.raw(n_value),
                increment_note_raw=c.raw(row.first("spt_increment_note")),
            )
        )

    def core_run(self, row: SheetRow, out: list[ExtractedCoreRun]) -> None:
        rqd, recovery = row.first("rqd"), row.first("recovery")
        if (rqd is None and recovery is None) or row.row_type not in ("layer", "core_run"):
            return
        c = self.cells
        out.append(
            ExtractedCoreRun(
                from_raw=c.raw(row.first("depth_from")),
                to_raw=c.raw(row.first("depth_to")),
                tcr_raw=c.raw(recovery),
                scr_raw=RawValue.absent(),
                rqd_raw=c.raw(rqd),
                ff_raw=RawValue.absent(),
            )
        )

    def groundwater(self, row: SheetRow, out: list[ExtractedGroundwater]) -> None:
        depth = row.first("groundwater_depth")
        if depth is None and row.row_type == "groundwater":
            depth = row.first("depth_from")
        on_date = row.first("groundwater_date")
        if depth is None and on_date is None:
            return
        c = self.cells
        out.append(
            ExtractedGroundwater(type=self._water_type(row), depth_raw=c.raw(depth), date_raw=c.raw(on_date))
        )

    def sample(self, row: SheetRow, out: list[ExtractedSample]) -> None:
        name = row.first("sample_name")
        if name is None or row.row_type not in ("sample", "layer"):
            return
        top, base = row.first("sample_from"), row.first("sample_to")
        if row.row_type == "sample":
            top = top or row.first("depth_from")
            base = base or row.first("depth_to")
        c = self.cells
        out.append(
            ExtractedSample(
                name_raw=c.raw(name),
                type_raw=c.raw(row.first("sample_type")),
                from_raw=c.raw(top),
                to_raw=c.raw(base),
            )
        )

    def tests(self, row: SheetRow, out: list[ExtractedTest]) -> None:
        if not row.tests:
            return
        depth = row.first("test_depth")
        sample = row.first("test_sample")
        if row.row_type == "test":
            depth = depth or row.first("depth_from")
            sample = sample or row.first("sample_name")
        c = self.cells
        for name, value, unit in row.tests:
            out.append(
                ExtractedTest(
                    name=name,
                    value_raw=c.raw(value),
                    unit_raw=c.raw(unit),
                    depth_raw=c.raw(depth),
                    sample_raw=c.raw(sample),
                )
            )


_SPT_FIELD = re.compile(r"^(stratigraphy\[[0-9]+\])\.spt_[a-z_]+$")


def _restore_raw(raw: str | None, texts: Mapping[str, str]) -> str | None:
    """The stored text of each cited cell instead of the text the parsers read ("4" not "4 ft")."""
    if raw is None or not texts:
        return raw
    return " / ".join(texts.get(part, part) for part in raw.split(" / "))


def _to_cell(entry: ProvenanceIn, cells: _Cells) -> ProvenanceIn:
    method = "read" if entry.method == "ocr" else entry.method
    return entry.model_copy(
        update={
            "page": None,
            "box": None,
            "cell": cells.ref_of_page.get(entry.page) if entry.page is not None else None,
            "method": method,
            "raw": _restore_raw(entry.raw, cells.text_of),
            "confidence": None,
        }
    )


def draft_borehole(document_id: str, bh: BoreholeRows) -> SheetDraft:
    """Normalise one borehole's rows like the PDF path and point every value at its cell (no database)."""
    result_builder = _Result(bh)
    result = result_builder.build()
    cells = result_builder.cells
    group = PageGroup(key=GROUP_PREFIX + bh.key, label_raw=bh.label.text, label_canonical=bh.key, pages=())
    outcome = ExtractionOutcome(result=result, blocks=dict(cells.blocks), run_id=None, model=None)
    draft: RecordDraft = build_record(document_id, group, outcome)
    warnings = [*draft.warnings]
    if draft.body is None:
        return SheetDraft(bh.key, bh.label.text, None, {}, "unknown", None, [], warnings)
    body = draft.body
    original = draft.provenance
    provenance = {path: _to_cell(entry, cells) for path, entry in original.items()}
    dropped: set[int] = set()
    # An SPT printed on a layer's row belongs to that layer; its own depth is not printed.
    for i, layer in enumerate(body.get("stratigraphy") or []):
        path = f"stratigraphy[{i}].spt_depth_m"
        entry = provenance.get(path)
        if entry is not None and entry.cell in result_builder.layer_spt_depths:
            layer["spt_depth_m"] = None
            dropped.add(id(original[path]))
            provenance.pop(path)
    # Values Python derived from a layer's printed SPT text (increments, refusal, penetration) cite that text.
    for path, entry in list(provenance.items()):
        match = _SPT_FIELD.match(path)
        if entry.cell is not None or match is None or path.endswith(".spt_raw"):
            continue
        source = provenance.get(f"{match.group(1)}.spt_raw")
        if source is not None and source.cell is not None:
            provenance[path] = entry.model_copy(
                update={
                    "cell": source.cell,
                    "raw": entry.raw or source.raw,
                    "method": "derived",
                    "state": "ok",
                }
            )
    # A DIN 4023 code printed in its own column is read, not derived from the description.
    for i, layer in enumerate(body.get("stratigraphy") or []):
        top = provenance.get(f"stratigraphy[{i}].depth_from_m")
        code_cell = result_builder.din_cells.get(top.cell or "") if top is not None else None
        if code_cell is not None and layer.get("din4023_code") == valid_din_code(code_cell.text):
            provenance[f"stratigraphy[{i}].din4023_code"] = ProvenanceIn(
                document_id=document_id, cell=code_cell.ref, raw=code_cell.text, method="read"
            )
    # Text values hold what the cell stores, not the parser's text.
    for path, entry in provenance.items():
        try:
            current = get_value(body, path)
        except PathError:
            continue
        if isinstance(current, str) and current in cells.text_of and entry.raw == cells.text_of[current]:
            set_value(body, path, cells.text_of[current])
    guessed = [_guessed(g, cells) for g in draft.guessed_depths if id(g.provenance) not in dropped]
    return SheetDraft(
        key=bh.key,
        label=bh.label.text,
        body=body,
        provenance=provenance,
        crs_status=draft.crs_status,
        crs_code=draft.crs_code,
        guessed=guessed,
        warnings=warnings,
    )


def _guessed(depth: GuessedDepth, cells: _Cells) -> tuple[str, str, float]:
    page = depth.provenance.page
    ref = cells.ref_of_page.get(page, "") if page is not None else ""
    return ref, cells.text_of.get(depth.raw, depth.raw), depth.value


# ----------------------------------------------------------------------------------------------
# Stages
# ----------------------------------------------------------------------------------------------


def _now() -> datetime:
    return datetime.now(UTC)


EMPTY_PROGRESS = {
    "pages_total": 0,
    "pages_classified": 0,
    "pages_ocr_total": 0,
    "pages_ocr_done": 0,
    "boreholes_extracted": 0,
}


async def _load_document(document_id: str) -> tuple[str, str, Path] | None:
    """(sha256, file name, path of the stored original) of a document, None when it is gone."""
    async with get_sessionmaker()() as session:
        doc = await session.get(Document, document_id)
        if doc is None:
            return None
        return doc.sha256, doc.file_name, get_storage().path(doc.storage_key)


async def _workbook(path: Path) -> Workbook:
    try:
        workbook = await asyncio.to_thread(read_workbook_cached, path)
    except SpreadsheetError as exc:
        raise StageError(str(exc), retryable=False) from exc
    if not any(sheet.cells for sheet in workbook.sheets):
        raise StageError("The spreadsheet has no filled cells.", retryable=False)
    return workbook


async def stage_sheet_read(document_id: str) -> None:
    """Open the file (no model); the document has no pages and page_count null."""
    loaded = await _load_document(document_id)
    if loaded is None:
        return
    workbook = await _workbook(loaded[2])
    async with get_sessionmaker()() as session:
        doc = await session.get(Document, document_id)
        if doc is None:
            return
        await session.execute(delete(Page).where(Page.document_id == document_id))
        doc.page_count = None
        doc.progress = dict(EMPTY_PROGRESS)
        doc.updated_at = _now()
        await session.commit()
    logger.info(
        "document %s: %s with %d sheet(s), %d filled cells",
        document_id,
        workbook.kind,
        len(workbook.sheets),
        workbook.cell_count,
    )


async def stage_sheet_map(document_id: str, mapper: SheetMapper | None = None) -> None:
    """The model's mapping of the file's table layouts (one cached call)."""
    loaded = await _load_document(document_id)
    if loaded is None:
        return
    workbook = await _workbook(loaded[2])
    preview = await asyncio.to_thread(build_preview, workbook)
    run = await ensure_mapping(document_id, loaded[0], preview, mapper=mapper)
    logger.info(
        "document %s: sheet mapping %s with %d region(s)%s",
        document_id,
        run.id,
        len(run.mapping.regions),
        " (stored)" if run.reused else "",
    )


async def _lock_document(session: AsyncSession, document_id: str) -> Document | None:
    return await session.scalar(
        select(Document)
        .where(Document.id == document_id)
        .with_for_update(key_share=True)
        .execution_options(populate_existing=True)
    )


async def _count_boreholes(session: AsyncSession, doc: Document) -> None:
    count = await session.scalar(
        select(func.count()).select_from(Borehole).where(Borehole.document_id == doc.id)
    )
    progress = {**EMPTY_PROGRESS, **(doc.progress or {})}
    progress["boreholes_extracted"] = int(count or 0)
    doc.progress = progress
    doc.updated_at = _now()


async def _ask_depth_unit(session: AsyncSession, doc: Document, uid: str, draft: SheetDraft) -> None:
    """One question per document: its depths have no printed unit and are stored as metres (guessed)."""
    key = f"unit_ambiguous:depth:{doc.id}"
    existing = await session.scalar(select(Clarification).where(Clarification.dedupe_key == key))
    if existing is not None:
        if existing.status != "answered" and uid not in (existing.borehole_uids or []):
            existing.borehole_uids = [*(existing.borehole_uids or []), uid]
            await session.flush()
            payload = clarification_out(existing).model_dump(mode="json")
            await emit(session, "clarification.created", payload, document_id=doc.id)
        return
    cell, text, metres = max(draft.guessed, key=lambda g: abs(g[2]))
    where = f" in cell {cell}" if cell else ""
    await create_clarification(
        session,
        kind="unit_ambiguous",
        document_id=doc.id,
        question=(
            f"The depths in {doc.file_name} have no printed unit (for example '{text}'{where}). They are "
            "stored as metres for now and marked as guessed. Which unit are they in?"
        ),
        fields=[
            {
                "name": "depth_unit",
                "widget": "number_unit",
                "label": "Depth unit",
                "help": f"Borehole {draft.label}: the value '{text}'{where}.",
                "required": True,
                "unit_options": list(DEPTH_UNIT_OPTIONS),
                "default": metres,
            }
        ],
        evidence=[],  # spreadsheet cells have no page to show
        borehole_uids=[uid],
        project=doc.project,
        dedupe_key=key,
    )


async def _create(document_id: str, draft: SheetDraft) -> str | None:
    """Create one borehole (idempotent per document and borehole id); returns its uid."""
    group_key = GROUP_PREFIX + draft.key
    async with get_sessionmaker()() as session:
        doc = await _lock_document(session, document_id)
        if doc is None or draft.body is None:
            return None
        existing = await session.scalar(
            select(Borehole.uid).where(Borehole.document_id == document_id, Borehole.group_key == group_key)
        )
        if existing is not None:
            return existing
        try:
            bh = await create_borehole(
                session,
                document=doc,
                group_key=group_key,
                body=draft.body,
                provenance=draft.provenance,
                crs_status=draft.crs_status,
                crs_code=draft.crs_code,
                pages=(),
                actor=ACTOR,
            )
        except IntegrityError:
            await session.rollback()
            return await session.scalar(
                select(Borehole.uid).where(
                    Borehole.document_id == document_id, Borehole.group_key == group_key
                )
            )
        if draft.guessed:
            await _ask_depth_unit(session, doc, bh.uid, draft)
        await _count_boreholes(session, doc)
        await emit_document_updated(session, doc)
        await session.commit()
        return bh.uid


async def stage_sheet_extract(document_id: str, mapper: SheetMapper | None = None) -> None:
    """Apply the stored mapping to every row, create one borehole per id and close the document."""
    loaded = await _load_document(document_id)
    if loaded is None:
        return
    workbook = await _workbook(loaded[2])
    preview = await asyncio.to_thread(build_preview, workbook)
    run = await ensure_mapping(document_id, loaded[0], preview, mapper=mapper)  # stored: no call
    drafts, warnings = await asyncio.to_thread(plan_drafts, document_id, run.mapping, workbook, preview)
    for warning in warnings:
        logger.warning("document %s: %s", document_id, warning)
    if drafts is None:
        raise StageError(
            "No table in this spreadsheet could be matched to borehole fields with a borehole id, so "
            "nothing was extracted.",
            retryable=False,
        )
    for draft in drafts:
        await _create(document_id, draft)
    await finalize_document(document_id)
    _LAST_READ.clear()  # the parsed file is not needed after the last stage


# Numbers checked against the model's decimal separator (a sample is enough to tell the style).
_STYLE_SAMPLE = 2000


def plan_drafts(
    document_id: str, mapping: SheetMapping, workbook: Workbook, preview: Preview
) -> tuple[list[SheetDraft] | None, list[str]]:
    """Validate the mapping, apply it to every row and draft every borehole (CPU only, no database).
    None when no region of the mapping is usable."""
    plan = validate_mapping(mapping, workbook, preview)
    warnings = [f"mapping: {w}" for w in plan.warnings]
    if not plan.regions:
        return None, warnings
    boreholes, row_warnings = apply_plan(workbook, plan)
    warnings += [f"rows: {w}" for w in row_warnings]
    if plan.decimal_separator is not None:
        texts = [c.parse for sh in workbook.sheets for c in sh.cells.values() if c.kind == "text"]
        learned = numbers.learn_style(texts[:_STYLE_SAMPLE]).decimal
        if learned is not None and learned != plan.decimal_separator:
            warnings.append(
                f"the model named '{plan.decimal_separator}' as decimal separator; the file's numbers use "
                f"'{learned}' (the file's style is used)"
            )
    drafts = []
    for bh in boreholes:
        draft = draft_borehole(document_id, bh)
        warnings += [f"borehole {bh.label.text}: {w}" for w in draft.warnings]
        drafts.append(draft)
    return drafts, warnings
