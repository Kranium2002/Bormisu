"""Stage 7: write results in the BoreholeDataFormat layout (xlsx) plus a JSON dump with full provenance."""
from __future__ import annotations

import json
import math
from dataclasses import asdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Borehole, METADATA_FIELDS, Page

META_COLS = METADATA_FIELDS + ["Latitude", "Longitude", "EPSG", "CRSStatus", "Client", "Contractor", "Method", "SourceFile", "Pages", "Confidence", "Issues"]
LOG_COLS = ["BoreholeID", "BoreholeName", "LayerNo", "DepthTop", "DepthBottom", "Thickness",
            "StratigraphicName", "EngineeringSoilUnit", "GeologicAge", "Lithology/SoilName", "Soil/Rockprofile",
            "Color", "Consistency", "Moisture", "Core", "SampleName", "SampleDepth", "TotalCoreRecovery", "RQD", "SPT",
            "Groundwaterlevel", "Description (raw)", "SourceFile", "Page", "y0", "y1", "Method", "Confidence",
            "NormConfidence"]
POINT_COLS = ["BoreholeID", "Kind", "Depth", "DepthTo", "Value", "Label", "SourceFile", "Page", "Method", "Confidence"]
PAGE_COLS = ["SourceFile", "Page", "Class", "Score", "TextLayer", "Words", "DepthScale", "ScaleTicks", "kwLog", "kwLab",
             "OCRconf", "Route", "ReadQuality", "ParserNotes"]


def _fmt(v: Any) -> Any:
    if isinstance(v, float) and math.isnan(v):
        return None
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    return v


def _in_interval(p, top: float, bottom: float) -> bool:
    if p.depth is None or (isinstance(p.depth, float) and math.isnan(p.depth)):
        return False
    d = p.depth if p.depth_to is None else (p.depth + p.depth_to) / 2
    return top - 1e-6 <= d < bottom + 1e-6


def _join(vals: list[Any]) -> Any:
    vals = [v for v in vals if v is not None]
    if not vals:
        return None
    if len(vals) == 1:
        return vals[0]
    return "; ".join(str(v) for v in vals)


def layer_rows(bh: Borehole) -> list[dict]:
    rows = []
    gw = bh.metadata.get("Groundwaterlevel")
    for i, L in enumerate(bh.layers, start=1):
        top, bot = L.top, L.bottom
        pts = [p for p in bh.points if not (math.isnan(top) or math.isnan(bot)) and _in_interval(p, top, bot)]
        spt = [p.value if p.value is not None else p.label for p in pts if p.kind == "SPT"]
        rqd = [p.value for p in pts if p.kind == "RQD"]
        tcr = [p.value for p in pts if p.kind == "TCR"]
        samples = [p for p in pts if p.kind == "SAMPLE"]
        core = "yes" if (rqd or tcr or (L.is_rock and "core" in L.description.lower())) else ("no" if L.is_rock is False else None)
        rows.append({
            "BoreholeID": bh.borehole_id, "BoreholeName": bh.metadata.get("BoreholeName"), "LayerNo": i,
            "DepthTop": _fmt(top), "DepthBottom": _fmt(bot),
            "Thickness": None if (math.isnan(top) or math.isnan(bot)) else round(bot - top, 2),
            "StratigraphicName": L.stratigraphic_name, "EngineeringSoilUnit": L.soil_unit, "GeologicAge": L.geologic_age,
            "Lithology/SoilName": L.lithology, "Soil/Rockprofile": L.din_code,
            "Color": L.color, "Consistency": L.consistency, "Moisture": L.moisture, "Core": core,
            "SampleName": _join([p.value or p.label for p in samples]),
            "SampleDepth": _join([p.depth if p.depth_to is None else f"{p.depth}-{p.depth_to}" for p in samples]),
            "TotalCoreRecovery": _join(tcr), "RQD": _join(rqd), "SPT": _join(spt),
            "Groundwaterlevel": gw if (isinstance(gw, (int, float)) and top <= gw < bot) else None,
            "Description (raw)": L.description, "SourceFile": Path(L.prov.file).name, "Page": L.prov.page,
            "y0": round(L.prov.y0, 1), "y1": round(L.prov.y1, 1), "Method": L.prov.method,
            "Confidence": round(L.prov.confidence, 2), "NormConfidence": L.norm_confidence,
        })
    return rows


def meta_row(bh: Borehole) -> dict:
    m = bh.metadata
    extra = m.get("_extra", {})
    row = {k: _fmt(m.get(k)) for k in METADATA_FIELDS}
    row["BoreholeID"] = bh.borehole_id
    row["Notes"] = _join([m.get("Notes"), json.dumps({k: v for k, v in extra.items() if k not in ("Client", "Contractor", "Method")}, ensure_ascii=False) if extra else None])
    row["Latitude"], row["Longitude"], row["EPSG"], row["CRSStatus"] = m.get("lat"), m.get("lon"), m.get("epsg"), m.get("crs_status")
    row["Client"] = extra.get("Client", m.get("Client"))
    row["Contractor"] = extra.get("Contractor", m.get("Contractor"))
    row["Method"] = extra.get("Method")
    row["SourceFile"] = Path(bh.source_file).name
    row["Pages"] = ",".join(str(p) for p in bh.pages)
    row["Confidence"] = m.get("_confidence")
    row["Issues"] = "; ".join(bh.issues) if bh.issues else None
    return row


def write_xlsx(boreholes: list[Borehole], pages: list[Page], page_notes: dict[str, dict[int, list[str]]], out: Path) -> Path:
    wb = Workbook()
    head_font = Font(bold=True)
    fill = PatternFill("solid", fgColor="DDEBF7")

    def sheet(title, cols, rows):
        ws = wb.create_sheet(title)
        ws.append(cols)
        for c in range(1, len(cols) + 1):
            ws.cell(row=1, column=c).font = head_font
            ws.cell(row=1, column=c).fill = fill
        for r in rows:
            ws.append([_fmt(r.get(c)) for c in cols])
        for i, c in enumerate(cols, start=1):
            width = max(len(str(c)), *(len(str(r.get(c) or "")) for r in rows)) if rows else len(c)
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(8, width + 2))
        ws.freeze_panes = "A2"
        return ws

    wb.remove(wb.active)
    sheet("Metadata", META_COLS, [meta_row(b) for b in boreholes])
    sheet("Logdata", LOG_COLS, [r for b in boreholes for r in layer_rows(b)])
    sheet("Points", POINT_COLS, [{
        "BoreholeID": b.borehole_id, "Kind": p.kind, "Depth": _fmt(p.depth), "DepthTo": p.depth_to, "Value": _fmt(p.value),
        "Label": p.label, "SourceFile": Path(p.prov.file).name, "Page": p.prov.page, "Method": p.prov.method,
        "Confidence": round(p.prov.confidence, 2)} for b in boreholes for p in b.points])
    sheet("Pages", PAGE_COLS, [{
        "SourceFile": Path(p.source_file).name, "Page": p.page_no, "Class": p.page_class, "Score": p.class_score,
        "TextLayer": p.has_text_layer, "Words": p.features.get("n_words"), "DepthScale": p.features.get("depth_scale_kind"),
        "ScaleTicks": p.features.get("depth_scale_n"), "kwLog": p.features.get("kw_log"), "kwLab": p.features.get("kw_lab"),
        "OCRconf": p.features.get("ocr_conf"), "Route": p.features.get("route"), "ReadQuality": p.features.get("read_quality"),
        "ParserNotes": "; ".join(page_notes.get(p.doc_id, {}).get(p.page_no, []))} for p in pages])
    sheet("Issues", ["BoreholeID", "Issue"], [{"BoreholeID": b.borehole_id, "Issue": i} for b in boreholes for i in b.issues])
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def write_json(boreholes: list[Borehole], out: Path) -> Path:
    def clean(o):
        if isinstance(o, float) and math.isnan(o):
            return None
        if isinstance(o, dict):
            return {k: clean(v) for k, v in o.items()}
        if isinstance(o, list):
            return [clean(v) for v in o]
        return o
    data = [clean(asdict(b)) for b in boreholes]
    out.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    return out
