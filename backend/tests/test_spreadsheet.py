"""Spreadsheet path (BACKEND.md 7.10): .xlsx and .csv uploads, unit by unit and end to end.

FICTIONAL data only: every workbook and CSV is generated here at test time; every label, place and value is
invented. The mapping model is a fake that returns a hand-written declarative mapping and records what it
was shown. No provider is called.

Three document families: one row per layer with a borehole id column (plus a location table), one sheet per
borehole with a header block, and a CSV with semicolons and decimal commas.
"""

from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Callable, Iterator, Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

import httpx
import openpyxl
import pytest
from openpyxl.comments import Comment
from sqlalchemy import func, select, update

from app.db.models import Borehole, Document, ExtractionRun
from app.db.session import get_sessionmaker
from app.pipeline import spreadsheet as sheets
from app.pipeline.orchestrator import run_pipeline
from app.pipeline.spreadsheet import (
    TARGET_SUPPORT,
    Preview,
    SheetMapping,
    Workbook,
    apply_plan,
    build_preview,
    column_index,
    config,
    number_texts,
    parse_ref,
    read_workbook,
    system_prompt,
    validate_mapping,
)
from tests.contract_utils import validate_response

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FT = 0.3048


# ----------------------------------------------------------------------------------------------
# Fictional files
# ----------------------------------------------------------------------------------------------


def xlsx_bytes(
    build: Callable[[openpyxl.Workbook], None], cached: Mapping[tuple[int, str], str] = {}
) -> bytes:
    """A workbook built by `build`; `cached` stores results for formula cells ((sheet number, cell) -> value)
    the way a spreadsheet program does (openpyxl itself writes formulas without results)."""
    book = openpyxl.Workbook()
    build(book)
    raw = io.BytesIO()
    book.save(raw)
    if not cached:
        return raw.getvalue()
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(raw.getvalue())) as src,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst,
    ):
        for info in src.infolist():
            content = src.read(info.filename)
            for (index, ref), value in cached.items():
                if info.filename == f"xl/worksheets/sheet{index}.xml":
                    pattern = re.compile(rf'(<c r="{ref}"[^>]*>\s*<f>[^<]*</f>)\s*<v\s*/>')
                    text, n = pattern.subn(rf"\g<1><v>{value}</v>", content.decode("utf-8"))
                    assert n == 1, ref
                    content = text.encode("utf-8")
            dst.writestr(info, content)
    return out.getvalue()


def layer_table_workbook() -> bytes:
    """Family 1: a location table and a log table with one row per layer and a borehole id column.

    SX-1's id is merged over its three rows; SX-2's id is printed on its first row only (blank below).
    Depths in feet (unit row), RQD in %, an SPT N column; one depth is a formula with a stored result and
    one RQD is a formula without one.
    """

    def build(book: openpyxl.Workbook) -> None:
        loc = book.active
        assert loc is not None
        loc.title = "Locations"
        loc["A1"] = "FICTIONAL SITE SX (test data, invented)"
        for col, title in zip(
            "ABCDE", ["Hole ID", "Easting", "Northing", "Grid", "Ground level (m)"], strict=True
        ):
            loc[f"{col}3"] = title
        loc.append(["SX-1", 413210.5, 5316044.2, "ETRS89 / UTM 32N", 312.4])
        loc.append(["SX-2", 413480.0, 5316310.7, None, 311.9])
        logs = book.create_sheet("Logs")
        logs["A1"] = "FICTIONAL BOREHOLE LOGS"
        logs.merge_cells("A1:F1")
        for col, title in zip(
            "ABCDEF", ["Hole ID", "From", "To", "Description", "SPT N", "RQD"], strict=True
        ):
            logs[f"{col}3"] = title
        logs["B3"].comment = Comment("Depths below ground level (fictional note)", "test")
        logs["B4"], logs["C4"], logs["F4"] = "ft", "ft", "%"
        rows: list[list[Any]] = [
            ["SX-1", 0, 4, "Sand, silty, brown", 12, None],
            [None, 4, "=B6+6", "Clay, stiff", None, None],
            [None, 10, 16.5, "Sandstone, weathered", "50/120", 85],
            ["SX-2", 0, 3.5, "Gravel, sandy", None, "=100*0.6"],
            [None, 3.5, 8, "Clay", None, None],
        ]
        for offset, row in enumerate(rows):
            for col, value in zip("ABCDEF", row, strict=True):
                if value is not None:
                    logs[f"{col}{5 + offset}"] = value
        logs.merge_cells("A5:A7")

    return xlsx_bytes(build, cached={(2, "C6"): "10"})


def layer_table_mapping() -> SheetMapping:
    return mapping(
        region(
            "Locations",
            3,
            [
                column("A", "borehole_id"),
                column("B", "x"),
                column("C", "y"),
                column("D", "coordinate_system"),
                column("E", "ground_level", unit="m", unit_cell="E3"),
            ],
            row_type="borehole",
        ),
        region(
            "Logs",
            3,
            [
                column("A", "borehole_id"),
                column("B", "depth_from", unit="ft", unit_cell="B4"),
                column("C", "depth_to", unit="ft", unit_cell="C4"),
                column("D", "description"),
                column("E", "spt_n"),
                column("F", "rqd", unit="%", unit_cell="F4"),
            ],
            unit_row=4,
            fill_down=["A"],
            language="en",
        ),
    )


HEADER_BLOCK = {
    "KB-7": {"easting": 512345.6, "northing": 5612345.7, "level": 45.2, "swap": False},
    "KB-8": {"easting": 512610.2, "northing": 5612590.4, "level": 44.8, "swap": True},
}


def sheet_per_borehole_workbook() -> bytes:
    """Family 2: one sheet per borehole with a header block above its log table. The second sheet prints
    northing before easting, so its values sit in other cells than the first sheet's."""

    def build(book: openpyxl.Workbook) -> None:
        first = book.active
        assert first is not None
        for index, (label, spec) in enumerate(HEADER_BLOCK.items()):
            ws = first if index == 0 else book.create_sheet()
            ws.title = "Log A" if index == 0 else "Log B"
            coords = [("Easting:", spec["easting"]), ("Northing:", spec["northing"])]
            if spec["swap"]:
                coords.reverse()
            block = [("Borehole:", label), *coords, ("Coordinate system:", "ETRS89 / UTM 33N")]
            block.append(("Ground level (m):", spec["level"]))
            for r, (name, value) in enumerate(block, start=1):
                ws[f"A{r}"], ws[f"B{r}"] = name, value
            for col, title in zip(
                "ABCD", ["Depth from [m]", "Depth to [m]", "Soil", "Water strike [m]"], strict=True
            ):
                ws[f"{col}7"] = title
            layers = [
                (0, 1.2, "Topsoil", None),
                (1.2, 3.75, "Sand, medium dense", 2.4),
                (3.75, 6, "Clay, stiff", None),
            ]
            for r, row in enumerate(layers, start=8):
                for col, value in zip("ABCD", row, strict=True):
                    if value is not None:
                        ws[f"{col}{r}"] = value

    return xlsx_bytes(build)


def sheet_per_borehole_mapping() -> SheetMapping:
    return mapping(
        region(
            "Log A",
            7,
            [
                column("A", "depth_from", unit="m", unit_cell="A7"),
                column("B", "depth_to", unit="m", unit_cell="B7"),
                column("C", "description"),
                column("D", "groundwater_depth", unit="m", unit_cell="D7"),
            ],
            [
                header_cell("A1", "B1", "borehole_id"),
                header_cell("A2", "B2", "x"),
                header_cell("A3", "B3", "y"),
                header_cell("A4", "B4", "coordinate_system"),
                header_cell("A5", "B5", "ground_level", unit="m", unit_cell="A5"),
            ],
            groundwater="strike",
            language="en",
        )
    )


CSV_TEXT = (
    "Bohrung;Tiefe von [m];Tiefe bis [m];Beschreibung;Rechtswert;Hochwert;Wassergehalt [%]\r\n"
    "KB-1;0,00;1,20;Mutterboden;3512345,6;5612345,7;\r\n"
    "KB-1;1,20;3,40;Sand, schluffig;3512345,6;5612345,7;18,5\r\n"
    "KB-1;3,40;5,00;Kies, sandig;3512345,6;5612345,7;\r\n"
    "KB-2;0,00;2,10;Löss, schluffig;3512500,1;5612400,9;21,0\r\n"
)


def csv_mapping() -> SheetMapping:
    return mapping(
        region(
            "csv",
            1,
            [
                column("A", "borehole_id"),
                column("B", "depth_from", unit="m", unit_cell="B1"),
                column("C", "depth_to", unit="m", unit_cell="C1"),
                column("D", "description"),
                column("E", "x"),
                column("F", "y"),
                column("G", "test_value", unit="%", unit_cell="G1", test_name="water content"),
            ],
            language="de",
        ),
        decimal=",",
    )


# ----------------------------------------------------------------------------------------------
# Mappings (what a model would answer) and the fake mapper
# ----------------------------------------------------------------------------------------------


def column(
    letter: str,
    target: str,
    *,
    unit: str | None = None,
    unit_cell: str | None = None,
    test_name: str | None = None,
) -> dict[str, Any]:
    return {"column": letter, "target": target, "test_name": test_name, "unit": unit, "unit_cell": unit_cell}


def header_cell(
    label: str | None, value: str, target: str, *, unit: str | None = None, unit_cell: str | None = None
) -> dict[str, Any]:
    return {"label_cell": label, "value_cell": value, "target": target, "unit": unit, "unit_cell": unit_cell}


def region(
    sheet: str,
    header_row: int | None,
    columns: Sequence[dict[str, Any]] = (),
    header_cells: Sequence[dict[str, Any]] = (),
    *,
    unit_row: int | None = None,
    row_type: str = "layer",
    fill_down: Sequence[str] = (),
    rules: Sequence[dict[str, Any]] = (),
    groundwater: str | None = None,
    language: str | None = None,
) -> dict[str, Any]:
    return {
        "sheet": sheet,
        "header_row": header_row,
        "unit_row": unit_row,
        "first_data_row": None,
        "last_data_row": None,
        "columns": list(columns),
        "header_cells": list(header_cells),
        "fill_down_columns": list(fill_down),
        "default_row_type": row_type,
        "row_type_rules": list(rules),
        "groundwater_type": groundwater,
        "description_language": language,
    }


def mapping(*regions: dict[str, Any], decimal: str | None = None) -> SheetMapping:
    return SheetMapping.model_validate({"regions": list(regions), "decimal_separator": decimal})


@dataclass
class FakeMapper:
    """The mapping model for tests: answers `mapping`, counts calls and keeps what it was shown."""

    mapping: SheetMapping
    calls: int = 0
    previews: list[str] = field(default_factory=list)

    async def map_sheets(
        self, preview: Preview, *, document_id: str
    ) -> tuple[SheetMapping, str | None, float]:
        self.calls += 1
        self.previews.append(preview.text)
        return self.mapping, "fake-mapper", 0.0


@pytest.fixture
def mapper(monkeypatch: pytest.MonkeyPatch) -> Iterator[FakeMapper]:
    fake = FakeMapper(mapping())
    monkeypatch.setattr(sheets, "OpenRouterSheetMapper", lambda: fake)
    yield fake


def workbook_of(tmp_path: Path, name: str, data: bytes) -> Workbook:
    path = tmp_path / name
    path.write_bytes(data)
    return read_workbook(path)


# ----------------------------------------------------------------------------------------------
# Unit tests (no database)
# ----------------------------------------------------------------------------------------------


def test_targets_and_prompt_come_from_the_data_files() -> None:
    cfg = config()
    assert set(cfg.targets) <= TARGET_SUPPORT and "borehole_id" in cfg.targets
    prompt = system_prompt()
    assert "{{" not in prompt and "test_value [row, test]" in prompt and "Watercontent [%]" in prompt
    assert "layer: One stratum per row" in prompt


def test_number_texts_and_cell_references() -> None:
    assert number_texts(1.234) == ("1.234", "1.2340")  # three decimals are never read as a thousands group
    assert number_texts(3.0) == ("3", "3") and number_texts(-0.5) == ("-0.5", "-0.5")
    assert number_texts(0.85, "0%") == ("85%", "85%")
    assert number_texts(0.1 + 0.2) == ("0.3", "0.3")
    assert parse_ref("$B$12") == (12, 2) and parse_ref("B0") is None and parse_ref("12") is None
    assert column_index("AA") == 27 and column_index("A1") is None


def test_xlsx_reader_keeps_merges_notes_and_stored_formula_results(tmp_path: Path) -> None:
    book = workbook_of(tmp_path, "fictional.xlsx", layer_table_workbook())
    assert [s.name for s in book.sheets] == ["Locations", "Logs"]
    logs = book.sheets[1]
    merged = logs.cell(6, 1)
    assert merged is not None and merged.text == "SX-1" and merged.source == (5, 1)  # merged id A5:A7
    assert book.ref(logs, merged.source) == "Logs!A5"
    cached = logs.cell(6, 3)
    assert cached is not None and (cached.text, cached.formula) == ("10", "=B6+6")
    assert logs.cell(8, 6) is None and logs.pending[(8, 6)] == "formula without a stored result"
    assert logs.notes[(3, 2)] == "Depths below ground level (fictional note)"
    title = logs.cell(1, 4)
    assert title is not None and title.merged_copy and title.text == "FICTIONAL BOREHOLE LOGS"
    assert "A1:F1" in logs.merged and "A5:A7" in logs.merged

    def odd_cells(b: openpyxl.Workbook) -> None:
        ws = b.active
        assert ws is not None
        ws.title = "Odd sheet"
        ws["A1"], ws["B1"], ws["C1"] = date(2025, 3, 5), 0.85, True
        ws["B1"].number_format = "0%"

    odd = workbook_of(tmp_path, "odd.xlsx", xlsx_bytes(odd_cells))
    sheet = odd.sheets[0]
    assert [(c.text, c.kind) for c in sheet.row_cells(1)] == [
        ("2025-03-05", "date"),
        ("85%", "number"),
        ("TRUE", "bool"),
    ]
    assert odd.ref(sheet, (1, 2)) == "'Odd sheet'!B1"


def test_csv_reader_sniffs_delimiter_and_encoding(tmp_path: Path) -> None:
    semicolons = workbook_of(tmp_path, "a.csv", CSV_TEXT.encode("cp1252"))
    sheet = semicolons.sheets[0]
    assert semicolons.kind == "csv" and sheet.max_col == 7 and sheet.max_row == 5
    description = sheet.cell(5, 4)
    assert description is not None and description.text == "Löss, schluffig"
    assert semicolons.ref(sheet, (3, 2)) == "B3"  # a CSV has no sheet name to cite
    commas = workbook_of(tmp_path, "b.csv", '﻿Hole,From,To,Description\nUX-1,0,1.5,"Sand, grey"\n'.encode())
    row = [c.text for c in commas.sheets[0].row_cells(2)]
    assert row == ["UX-1", "0", "1.5", "Sand, grey"]
    header = commas.sheets[0].cell(1, 1)
    assert header is not None and header.text == "Hole"  # the byte-order mark is not part of the text


def long_table_workbook(rows: int = 60) -> bytes:
    def build(book: openpyxl.Workbook) -> None:
        ws = book.active
        assert ws is not None
        ws.title = "Data"
        ws["A1"] = "Fictional data table"
        ws.append([])
        ws.append(["Hole", "Top", "Base", "Description", "Kind"])
        for i in range(rows):
            kind = "SPT" if i % 10 == 9 else "L"
            ws.append([f"LT-{i // 20 + 1}", i * 0.5, i * 0.5 + 0.5, f"Fictional layer {i + 1}", kind])

    return xlsx_bytes(build)


def test_preview_holds_headers_and_a_few_rows_never_the_table(tmp_path: Path) -> None:
    book = workbook_of(tmp_path, "long.xlsx", long_table_workbook())
    preview = build_preview(book)
    text = preview.text
    assert 'R3: A3="Hole" | B3="Top" | C3="Base" | D3="Description" | E3="Kind"' in text
    assert "detected table: header row 3, data rows 4-63 (60 rows with values)" in text
    assert "Fictional layer 5" in text and "Fictional layer 6" not in text  # 5 sample rows only
    assert "Fictional layer 40" not in text
    assert "(55 more rows with values are not shown)" in text
    assert 'column E values: "L" x54, "SPT" x6' in text  # lets the model write a row-type rule
    assert len(text) < 1500


def test_preview_shows_a_further_table_lower_on_the_sheet(tmp_path: Path) -> None:
    def build(book: openpyxl.Workbook) -> None:
        ws = book.active
        assert ws is not None
        ws.append(["Hole", "Top", "Base", "Description"])
        for i in range(12):
            ws.append(["FT-1", i, i + 1, f"Fictional stratum {i + 1}"])
        ws.append([])
        ws.append(["Hole", "Sample", "Depth", "Water content"])
        for i in range(4):
            ws.append(["FT-1", f"S{i + 1}", i + 0.5, 20 + i])

    preview = build_preview(workbook_of(tmp_path, "two-tables.xlsx", xlsx_bytes(build)))
    assert 'R15: A15="Hole" | B15="Sample" | C15="Depth" | D15="Water content"' in preview.text
    assert 'R16: A16="FT-1" | B16="S1" | C16=0.5 | D16=20' in preview.text
    assert "Fictional stratum 9" not in preview.text  # rows between the two tables stay hidden


def test_layouts_share_one_preview_section(tmp_path: Path) -> None:
    book = workbook_of(tmp_path, "family2.xlsx", sheet_per_borehole_workbook())
    preview = build_preview(book)
    assert [layout.names for layout in preview.layouts] == [("Log A", "Log B")]
    assert 'same layout (this sheet\'s mapping applies to them too): "Log B"' in preview.text
    assert "KB-7" in preview.text and "KB-8" not in preview.text


def test_validation_keeps_what_exists_and_names_the_rest(tmp_path: Path) -> None:
    book = workbook_of(tmp_path, "family1.xlsx", layer_table_workbook())
    preview = build_preview(book)
    bad = mapping(
        region("No such sheet", 1, [column("A", "borehole_id")]),
        region("Locations", 3, [column("B", "x")]),  # no borehole id anywhere
        region(
            "Logs",
            3,
            [
                column("A", "borehole_id"),
                column("B", "depth_from", unit="m", unit_cell="B4"),  # the unit row prints ft
                column("C", "depth_to"),
                column("D", "lithology_code"),  # not a target
                column("E", "test_value", test_name="shoe size"),  # not a Test Results name
                column("Z", "description"),  # beyond the sheet
                column("F", "rqd", unit="%", unit_cell="F4"),
                column("C", "colour"),  # a second target for one column
            ],
            unit_row=4,
        ),
    )
    plan = validate_mapping(bad, book, preview)
    assert [r.sheets for r in plan.regions] == [("Logs",)]
    (logs,) = plan.regions
    assert {c.target: c.col for c in logs.columns} == {
        "borehole_id": 1,
        "depth_from": 2,
        "depth_to": 3,
        "rqd": 6,
    }
    units_used = {c.target: c.unit.key if c.unit else None for c in logs.columns}
    assert units_used == {"borehole_id": None, "depth_from": "ft", "depth_to": "ft", "rqd": "%"}
    joined = "\n".join(plan.warnings)
    for part in (
        "No such sheet",
        "no borehole id column or cell",
        "'lithology_code' is not a target field",
        "'shoe size' is not a Test Results name",
        "column Z: no such column",
        "the model named unit 'm'; the printed 'ft' is used",
        "already mapped to another target",
    ):
        assert part in joined, part


def test_rows_are_grouped_by_id_with_merges_and_fill_down(tmp_path: Path) -> None:
    book = workbook_of(tmp_path, "family1.xlsx", layer_table_workbook())
    preview = build_preview(book)
    plan = validate_mapping(layer_table_mapping(), book, preview)
    boreholes, warnings = apply_plan(book, plan)
    assert [b.label.text for b in boreholes] == ["SX-1", "SX-2"]
    sx1, sx2 = boreholes
    assert [r.row for r in sx1.rows] == [5, 6, 7] and [r.row for r in sx2.rows] == [8, 9]
    assert sx1.label.ref == "Locations!A4" and sx1.header["x"].ref == "Locations!B4"
    assert "coordinate_system" not in sx2.header  # blank for SX-2: never filled in
    assert [v.parse for v in sx1.rows[0].values["depth_to"]] == ["4 ft"]  # the printed unit is attached
    assert sx1.rows[0].values["depth_to"][0].text == "4"
    assert (
        "sheet 'Logs' column F: 1 cell(s) with formula without a stored result not read (Logs!F8)" in warnings
    )


def target_fields_workbook() -> bytes:
    """A sheet per borehole whose header block prints a name, a project number and a length along the hole
    in feet, above a log table with an engineering-unit and an age column (fictional values)."""

    def build(book: openpyxl.Workbook) -> None:
        ws = book.active
        assert ws is not None
        ws.title = "LogT"
        block = [
            ("Hole:", "TX-1"),
            ("Name:", "Fictional east well"),
            ("Project No.:", "P-881"),
            ("Length along hole (ft):", 32.5),
        ]
        for r, (name, value) in enumerate(block, start=1):
            ws[f"A{r}"], ws[f"B{r}"] = name, value
        for col, title in zip("ABCDE", ["From [m]", "To [m]", "Soil", "Unit", "Age"], strict=True):
            ws[f"{col}6"] = title
        rows = [(0, 2.5, "Clay, soft", "U1", "Holocene"), (2.5, 9.5, "Marl, hard", "U2", None)]
        for r, row in enumerate(rows, start=7):
            for col, value in zip("ABCDE", row, strict=True):
                if value is not None:
                    ws[f"{col}{r}"] = value

    return xlsx_bytes(build)


def test_target_fields_are_read_from_their_cells(tmp_path: Path) -> None:
    book = workbook_of(tmp_path, "family-targets.xlsx", target_fields_workbook())
    preview = build_preview(book)
    plan = mapping(
        region(
            "LogT",
            6,
            [
                column("A", "depth_from", unit="m", unit_cell="A6"),
                column("B", "depth_to", unit="m", unit_cell="B6"),
                column("C", "description"),
                column("D", "engineering_soil_unit"),
                column("E", "geologic_age"),
            ],
            [
                header_cell("A1", "B1", "borehole_id"),
                header_cell("A2", "B2", "borehole_name"),
                header_cell("A3", "B3", "project_id"),
                header_cell("A4", "B4", "borehole_length", unit="ft", unit_cell="A4"),
            ],
            language="en",
        )
    )
    drafts, warnings = sheets.plan_drafts("doc_X", plan, book, preview)
    assert drafts is not None and len(drafts) == 1, warnings
    draft = drafts[0]
    assert draft.body is not None
    meta, prov = draft.body["metadata"], draft.provenance
    assert (meta["borehole_id"], meta["borehole_name"], meta["project_id"]) == (
        "TX-1",
        "Fictional east well",
        "P-881",
    )
    assert meta["project_name"] is None  # a number is not a name
    # The length carries its own printed unit (ft) and is converted; the cell keeps the stored text.
    assert meta["borehole_length_m"] == pytest.approx(32.5 * FT)
    length = prov["metadata.borehole_length_m"]
    assert (length.cell, length.raw, length.method) == ("LogT!B4", "32.5", "read")
    assert meta["borehole_depth_m"] == 9.5  # the depth stays the depth (last layer)
    layers = draft.body["stratigraphy"]
    assert [(la["engineering_soil_unit"], la["geologic_age"]) for la in layers] == [
        ("U1", "Holocene"),
        ("U2", None),
    ]
    assert (
        prov["stratigraphy[1].engineering_soil_unit"].cell,
        prov["stratigraphy[0].geologic_age"].cell,
    ) == (
        "LogT!D8",
        "LogT!E7",
    )
    assert "stratigraphy[1].geologic_age" not in prov  # blank cell: null, no entry
    assert draft.guessed == []
    prompt = system_prompt()
    for target in ("borehole_name", "project_id", "borehole_length", "engineering_soil_unit", "geologic_age"):
        assert f"\n{target} [" in prompt, target


# ----------------------------------------------------------------------------------------------
# Upload to review (database)
# ----------------------------------------------------------------------------------------------


async def call(client: httpx.AsyncClient, method: str, path: str, *, expect: int, **kwargs: Any) -> Any:
    response = await client.request(method, path, **kwargs)
    assert response.status_code == expect, response.text
    content_type = response.headers.get("content-type", "application/json").split(";")[0]
    body = response.json() if response.content and "json" in content_type else None
    validate_response(method, path.split("?")[0], response.status_code, body, content_type)
    return body


async def upload(
    client: httpx.AsyncClient, name: str, data: bytes, mime: str, project: str | None = None
) -> str:
    form = {"project": project} if project else {}
    body = await call(
        client, "POST", "/api/v1/documents", expect=202, files={"file": (name, data, mime)}, data=form
    )
    return str(body["document_id"])


async def work(queue: Any, document_id: str) -> None:
    jobs = [
        j for j in queue.jobs.values() if j["status"] == "todo" and j["args"] == {"document_id": document_id}
    ]
    assert [j["task_name"] for j in jobs] == ["process_document"]
    jobs[0]["status"] = "doing"
    await run_pipeline(document_id)
    await queue.finish_job_run(jobs[0]["id"], status="succeeded", delete_job=False)


async def document(client: httpx.AsyncClient, document_id: str) -> dict[str, Any]:
    body: dict[str, Any] = await call(client, "GET", f"/api/v1/documents/{document_id}", expect=200)
    return body


async def records(client: httpx.AsyncClient, document_id: str) -> dict[str, dict[str, Any]]:
    body = await call(client, "GET", f"/api/v1/documents/{document_id}/boreholes", expect=200)
    return {r["metadata"]["borehole_id"]: r for r in body}


async def count(model: Any, document_id: str) -> int:
    async with get_sessionmaker()() as session:
        value = await session.scalar(
            select(func.count()).select_from(model).where(model.document_id == document_id)
        )
    return int(value or 0)


def assert_cell(entry: dict[str, Any], cell: str, raw: str | None, method: str = "read") -> None:
    assert (entry["cell"], entry["raw"], entry["method"]) == (cell, raw, method), entry
    assert entry["page"] is None and entry["box"] is None, entry


def assert_cells_only(record: dict[str, Any], document_id: str) -> None:
    """Every value of a spreadsheet record points at a cell, never at a page."""
    for path, entry in record["provenance"].items():
        assert entry["document_id"] == document_id, path
        assert entry["page"] is None and entry["box"] is None, path
        assert entry["cell"], path
        assert entry["method"] in ("read", "derived", "guessed"), (path, entry["method"])


@pytest.mark.db
async def test_layer_table_workbook_end_to_end(
    client: httpx.AsyncClient, live_mode: None, queue: Any, mapper: FakeMapper
) -> None:
    mapper.mapping = layer_table_mapping()
    doc_id = await upload(
        client, "fictional-logs.xlsx", layer_table_workbook(), XLSX_MIME, project="Fictional SX"
    )
    await work(queue, doc_id)

    doc = await document(client, doc_id)
    assert doc["status"] == "needs_review" and doc["error"] is None
    assert doc["page_count"] is None and doc["pages"] == []
    assert doc["progress"]["boreholes_extracted"] == 2
    assert doc["found"] == {"boreholes": 2, "layers": 5, "spt": 2, "groundwater": 0}
    assert mapper.calls == 1
    shown = mapper.previews[0]
    assert "Depths below ground level (fictional note)" in shown and 'A8="SX-2"' in shown

    by_label = await records(client, doc_id)
    assert sorted(by_label) == ["SX-1", "SX-2"]
    sx1, sx2 = by_label["SX-1"], by_label["SX-2"]
    for record in (sx1, sx2):
        assert_cells_only(record, doc_id)
        assert record["pages"] == [] and record["source_file"] == "fictional-logs.xlsx"
        assert record["groundwater"] == [] and record["samples"] == [] and record["test_results"] == []

    # Feet converted to metres; the record keeps the cell and the printed text.
    layers = sx1["stratigraphy"]
    assert [(s["depth_from_m"], s["depth_to_m"]) for s in layers] == [
        (0.0, pytest.approx(4 * FT)),
        (pytest.approx(4 * FT), pytest.approx(10 * FT)),
        (pytest.approx(10 * FT), pytest.approx(16.5 * FT)),
    ]
    prov = sx1["provenance"]
    assert_cell(prov["stratigraphy[0].depth_to_m"], "Logs!C5", "4")
    assert_cell(prov["stratigraphy[1].depth_to_m"], "Logs!C6", "10")  # a formula's stored result
    assert_cell(prov["stratigraphy[2].original_description"], "Logs!D7", "Sandstone, weathered")
    assert_cell(prov["metadata.borehole_id"], "Locations!A4", "SX-1")
    assert [s["original_description"] for s in layers] == [
        "Sand, silty, brown",
        "Clay, stiff",
        "Sandstone, weathered",
    ]
    assert [s["din4023_code"] for s in layers] == ["S", "T", "Sst"]
    assert prov["stratigraphy[0].din4023_code"]["method"] == "derived"
    assert all(s["english_lithology"] is None for s in layers)  # never invented

    # SPT on a layer's row belongs to that layer; its own depth is not printed, so it stays null.
    assert (layers[0]["spt_n_value"], layers[0]["spt_raw"], layers[0]["spt_depth_m"]) == (12, "12", None)
    assert "stratigraphy[0].spt_depth_m" not in prov
    assert_cell(prov["stratigraphy[0].spt_raw"], "Logs!E5", "12")
    assert layers[2]["spt_refusal"] is True and layers[2]["spt_n_value"] is None
    assert layers[1]["spt_n_value"] is None and layers[1]["spt_raw"] is None

    # RQD in %, from the row of its layer.
    assert (layers[2]["rqd_pct"], layers[2]["is_core"]) == (85.0, True)
    assert_cell(prov["stratigraphy[2].rqd_pct"], "Logs!F7", "85")
    assert prov["stratigraphy[2].is_core"]["method"] == "derived"
    assert layers[0]["rqd_pct"] is None and layers[0]["is_core"] is None

    # Total depth from the last layer; ground level in metres from the location table.
    meta = sx1["metadata"]
    assert (
        meta["borehole_depth_m"] == pytest.approx(16.5 * FT) and meta["borehole_depth_source"] == "last_layer"
    )
    assert meta["surface_elevation_mASL"] == 312.4
    assert_cell(prov["metadata.surface_elevation_mASL"], "Locations!E4", "312.4")

    # A stated system plots; a missing one is asked once and never filled in.
    coords = meta["coordinates"]
    assert (coords["local_x"], coords["local_y"]) == (413210.5, 5316044.2)
    assert coords["x_raw"] == "413210.5" and coords["axis_labels_raw"] == "Easting, Northing"
    assert coords["coordinate_system"] == "ETRS89 / UTM 32N" and coords["crs_code"] == "EPSG:25832"
    assert sx1["crs_status"] == "stated" and sx1["plotted"] is True
    assert_cell(prov["metadata.coordinates.x_raw"], "Locations!B4", "413210.5")
    assert prov["metadata.coordinates.latitude"]["cell"] == "Locations!B4"
    assert sx2["crs_status"] == "inferred" and sx2["plotted"] is False
    assert sx2["unplotted_reason"] == "crs_unconfirmed"
    assert sx2["metadata"]["coordinates"]["coordinate_system"] is None

    # SX-2: the id printed once covers the blank cell below; an RQD formula without a stored result is not
    # read and nothing replaces it.
    assert [(s["depth_from_m"], s["depth_to_m"]) for s in sx2["stratigraphy"]] == [
        (0.0, pytest.approx(3.5 * FT)),
        (pytest.approx(3.5 * FT), pytest.approx(8 * FT)),
    ]
    assert_cell(sx2["provenance"]["stratigraphy[1].depth_from_m"], "Logs!B9", "3.5")
    assert all(s["rqd_pct"] is None for s in sx2["stratigraphy"])

    questions = await call(client, "GET", "/api/v1/clarifications", expect=200)
    crs = [q for q in questions if q["kind"] == "crs_confirm"]
    assert len(crs) == 1 and crs[0]["borehole_uids"] == [sx2["uid"]]

    # Idempotent: a rerun (whole pipeline, or from the mapping stage) pays nothing and creates nothing.
    await run_pipeline(doc_id)
    async with get_sessionmaker()() as session:
        await session.execute(update(Document).where(Document.id == doc_id).values(stage="sheet_read"))
        await session.commit()
    await run_pipeline(doc_id)
    assert mapper.calls == 1
    assert await count(Borehole, doc_id) == 2 and await count(ExtractionRun, doc_id) == 1
    again = await records(client, doc_id)
    assert {k: r["revision"] for k, r in again.items()} == {"SX-1": 1, "SX-2": 1}
    assert (await document(client, doc_id))["status"] == "needs_review"

    # Review works on cell-cited values like on page-cited ones.
    values = f"/api/v1/boreholes/{sx1['uid']}/values"
    edit = {"path": "stratigraphy[1].original_description", "action": "edit", "value": "Clay, stiff, grey"}
    edited = await call(client, "PATCH", values, expect=200, json=edit)
    entry = edited["provenance"]["stratigraphy[1].original_description"]
    assert entry["method"] == "edited" and entry["cell"] == "Logs!D6" and entry["ai_value"] == "Clay, stiff"


@pytest.mark.db
async def test_sheet_per_borehole_workbook(
    client: httpx.AsyncClient, live_mode: None, queue: Any, mapper: FakeMapper
) -> None:
    mapper.mapping = sheet_per_borehole_mapping()
    doc_id = await upload(client, "fictional-sheets.xlsx", sheet_per_borehole_workbook(), XLSX_MIME)
    await work(queue, doc_id)

    doc = await document(client, doc_id)
    assert doc["status"] == "needs_review" and doc["page_count"] is None
    assert mapper.calls == 1 and "KB-8" not in mapper.previews[0]  # one section for both sheets

    by_label = await records(client, doc_id)
    assert sorted(by_label) == ["KB-7", "KB-8"]
    for label_text, record in by_label.items():
        spec = HEADER_BLOCK[label_text]
        sheet_ref = "'Log A'" if label_text == "KB-7" else "'Log B'"
        assert_cells_only(record, doc_id)
        coords = record["metadata"]["coordinates"]
        assert (coords["local_x"], coords["local_y"]) == (spec["easting"], spec["northing"])
        assert record["crs_status"] == "stated" and coords["crs_code"] == "EPSG:25833" and record["plotted"]
        x_cell = "B3" if spec["swap"] else "B2"  # found by its label on the sheet that prints it lower
        assert_cell(
            record["provenance"]["metadata.coordinates.x_raw"], f"{sheet_ref}!{x_cell}", str(spec["easting"])
        )
        assert record["metadata"]["surface_elevation_mASL"] == spec["level"]
        assert [(s["depth_from_m"], s["depth_to_m"]) for s in record["stratigraphy"]] == [
            (0.0, 1.2),
            (1.2, 3.75),
            (3.75, 6.0),
        ]
        assert_cell(record["provenance"]["stratigraphy[1].depth_to_m"], f"{sheet_ref}!B9", "3.75")
        (water,) = record["groundwater"]
        assert (water["type"], water["depth_m"], water["raw"]) == ("strike", 2.4, "2.4")
        assert_cell(record["provenance"]["groundwater[0].depth_m"], f"{sheet_ref}!D9", "2.4")
        assert record["provenance"]["groundwater[0].type"]["method"] == "guessed"
        assert record["metadata"]["borehole_depth_m"] == 6.0
    questions = await call(client, "GET", "/api/v1/clarifications", expect=200)
    assert [q for q in questions if q["kind"] == "crs_confirm"] == []


@pytest.mark.db
async def test_csv_with_semicolons_and_decimal_commas(
    client: httpx.AsyncClient, live_mode: None, queue: Any, mapper: FakeMapper
) -> None:
    mapper.mapping = csv_mapping()
    doc_id = await upload(client, "fictional-bohrungen.csv", CSV_TEXT.encode("cp1252"), "text/csv")
    await work(queue, doc_id)

    doc = await document(client, doc_id)
    assert doc["status"] == "needs_review" and doc["page_count"] is None and doc["pages"] == []
    by_label = await records(client, doc_id)
    assert sorted(by_label) == ["KB-1", "KB-2"]
    kb1, kb2 = by_label["KB-1"], by_label["KB-2"]
    for record in (kb1, kb2):
        assert_cells_only(record, doc_id)
    assert [(s["depth_from_m"], s["depth_to_m"]) for s in kb1["stratigraphy"]] == [
        (0.0, 1.2),
        (1.2, 3.4),
        (3.4, 5.0),
    ]
    assert [s["din4023_code"] for s in kb1["stratigraphy"]] == ["Mu", "S", "G"]
    assert_cell(kb1["provenance"]["stratigraphy[0].depth_to_m"], "C2", "1,20")
    assert kb2["stratigraphy"][0]["original_description"] == "Löss, schluffig"
    coords = kb1["metadata"]["coordinates"]
    assert (coords["x_raw"], coords["local_x"], coords["local_y"]) == ("3512345,6", 3512345.6, 5612345.7)
    assert coords["axis_labels_raw"] == "Rechtswert, Hochwert"
    assert kb1["crs_status"] == "inferred" and kb1["plotted"] is False
    (test,) = kb1["test_results"]
    assert (test["name"], test["value"], test["unit"], test["depth_m"]) == ("Watercontent", 18.5, "%", None)
    assert_cell(kb1["provenance"]["test_results[0].value"], "G3", "18,5 %")  # value and its column's unit
    assert kb2["test_results"][0]["value"] == 21.0
    questions = await call(client, "GET", "/api/v1/clarifications", expect=200)
    (crs,) = [q for q in questions if q["kind"] == "crs_confirm"]
    assert sorted(crs["borehole_uids"]) == sorted([kb1["uid"], kb2["uid"]])


@pytest.mark.db
async def test_depths_without_a_printed_unit_are_guessed_and_asked_once(
    client: httpx.AsyncClient, live_mode: None, queue: Any, mapper: FakeMapper
) -> None:
    mapper.mapping = mapping(
        region(
            "csv",
            1,
            [
                column("A", "borehole_id"),
                column("B", "depth_from"),
                column("C", "depth_to"),
                column("D", "description"),
            ],
            language="en",
        )
    )
    text = "Hole,From,To,Description\nUX-1,0,1.5,Sand\nUX-1,1.5,3.25,Clay\nUX-2,0,2,Gravel\n"
    doc_id = await upload(client, "fictional-no-units.csv", text.encode(), "text/csv")
    await work(queue, doc_id)

    by_label = await records(client, doc_id)
    ux1 = by_label["UX-1"]
    assert [(s["depth_from_m"], s["depth_to_m"]) for s in ux1["stratigraphy"]] == [(0.0, 1.5), (1.5, 3.25)]
    assert_cell(ux1["provenance"]["stratigraphy[1].depth_to_m"], "C3", "3.25", method="guessed")
    questions = await call(client, "GET", "/api/v1/clarifications", expect=200)
    (unit,) = [q for q in questions if q["kind"] == "unit_ambiguous"]
    assert sorted(unit["borehole_uids"]) == sorted(r["uid"] for r in by_label.values())
    assert unit["evidence"] == [] and "cell C3" in unit["question"]
    assert unit["fields"][0]["unit_options"] == ["m", "ft"]


@pytest.mark.db
async def test_a_mapping_without_usable_regions_fails_with_a_reason(
    client: httpx.AsyncClient, live_mode: None, queue: Any, mapper: FakeMapper
) -> None:
    mapper.mapping = mapping(region("Elsewhere", 1, [column("A", "borehole_id")]))
    doc_id = await upload(client, "fictional-unmapped.xlsx", long_table_workbook(10), XLSX_MIME)
    await work(queue, doc_id)
    doc = await document(client, doc_id)
    assert doc["status"] == "failed" and doc["page_count"] is None
    assert doc["error"] == (
        "No table in this spreadsheet could be matched to borehole fields with a borehole id, so nothing "
        "was extracted."
    )
    assert await count(Borehole, doc_id) == 0 and mapper.calls == 1
